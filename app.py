import os, base64, urllib.parse, csv, io, logging, secrets
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta, timezone
from typing import List, Dict, Any, Optional, Tuple
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
import psycopg2
from psycopg2.extras import RealDictCursor, execute_values
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

load_dotenv()
logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

st.set_page_config(page_title="Qween | Acompanhamento da Execução", page_icon="📊", layout="wide", initial_sidebar_state="collapsed")

ADO_ORG = "SME-Spassu"
ADO_BASE_URL = f"https://dev.azure.com/{ADO_ORG}"
ALL_WI_TYPES = ["Task", "Product Backlog Item", "Melhoria", "HotFix", "Design", "Bug", "BugFix", "Débito Técnico"]
ALL_STATES = ["To Do", "In Progress", "Resolved", "Done"]
DEFAULT_WI_TYPES = [t for t in ALL_WI_TYPES if t not in ["Task", "Débito Técnico"]]

SPASSU_LOGO_URL = "https://raw.githubusercontent.com/marcuscaldevilla-lang/sme-dashboard/main/spassu4.png"
SME_LOGO_URL = "https://raw.githubusercontent.com/marcuscaldevilla-lang/sme-dashboard/main/sme.png"

# =============================================================================
# PERFORMANCE: SESSÃO HTTP REUTILIZÁVEL COM POOL
# =============================================================================
@st.cache_resource
def get_http_session() -> requests.Session:
    session = requests.Session()
    retry_strategy = Retry(total=3, backoff_factor=1, status_forcelist=[429, 500, 502, 503, 504])
    adapter = HTTPAdapter(max_retries=retry_strategy, pool_connections=25, pool_maxsize=25)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    return session

def inject_custom_css():
    css = """
    <style>
        #MainMenu {visibility: hidden;} footer {visibility: hidden;} header {visibility: hidden;}
        .stApp { background: radial-gradient(circle at 15% 45%, #ff5d2b 0%, #f24e1e 50%, #df3a09 100%) !important; font-family: 'Montserrat', sans-serif; }
        .dotted-rings { position: fixed; top: 50%; left: -150px; transform: translateY(-50%); width: 900px; height: 900px; pointer-events: none; z-index: 0; }
        .ring { position: absolute; top: 50%; left: 50%; border-radius: 50%; border: 3px dotted rgba(255, 255, 255, 0.85); transform: translate(-50%, -50%); }
        .ring-1 { width: 300px; height: 300px; } .ring-2 { width: 440px; height: 440px; border-width: 4px; }
        .ring-3 { width: 580px; height: 580px; border-width: 4px; } .ring-4 { width: 720px; height: 720px; border-width: 5px; opacity: 0.9; }
        .ring-5 { width: 880px; height: 880px; border-width: 5px; opacity: 0.7; }
        .stTextInput > div > div > input, .stNumberInput > div > div > input, .stSelectbox > div > div > div > input { background-color: rgba(255, 255, 255, 0.95) !important; color: #222222 !important; border-radius: 10px !important; height: 48px !important; font-weight: 500 !important; }
        .stButton > button { width: 100%; height: 50px; background-color: #ffffff !important; color: #f24e1e !important; font-weight: 700 !important; border-radius: 10px !important; border: none !important; }
        .hero-title { color: #ffffff; font-size: 3.2rem; font-weight: 800; line-height: 1.15; margin-bottom: 20px; text-shadow: 0 4px 14px rgba(0, 0, 0, 0.15); }
        .hero-subtitle { color: #e6e6e6; font-size: 1.25rem; font-weight: 500; }
        .badge-qween { display: inline-block; background: rgba(255, 255, 255, 0.22); backdrop-filter: blur(8px); padding: 6px 16px; border-radius: 30px; font-size: 0.85rem; font-weight: 700; letter-spacing: 1.5px; text-transform: uppercase; margin-bottom: 20px; color: #ffffff; border: 1px solid rgba(255, 255, 255, 0.35); }
        .auth-card { background: rgba(255, 255, 255, 0.12); backdrop-filter: blur(20px); -webkit-backdrop-filter: blur(20px); border: 1px solid rgba(255, 255, 255, 0.25); border-radius: 20px; padding: 40px; box-shadow: 0 20px 45px rgba(0, 0, 0, 0.18); color: #ffffff; }
        .token-box { background: rgba(0,0,0,0.4); padding: 15px; border-radius: 8px; font-family: monospace; font-size: 1rem; word-break: break-all; margin: 15px 0; border: 2px solid rgba(255,255,255,0.5); text-align: center; color: #fff; }
    </style>
    <link href="https://fonts.googleapis.com/css2?family=Montserrat:wght@400;500;600;700;800&display=swap" rel="stylesheet">
    """
    st.markdown(css, unsafe_allow_html=True)

# =============================================================================
# AUTENTICAÇÃO
# =============================================================================
def authenticate_user(email: str, password: str) -> Optional[Dict]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT id, nome, email, tipo, apikey, status FROM usuarios WHERE email = %s AND senha_hash = %s LIMIT 1", (email.lower().strip(), password))
            user = cur.fetchone()
        conn.close()
        return user
    except Exception as e:
        logger.error(f"Erro na autenticação: {e}")
        return None

def criar_usuario_com_apikey(nome: str, email: str, senha: str, apikey: str) -> Tuple[bool, str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado."
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor() as cur:
            cur.execute("SELECT id FROM usuarios WHERE email = %s", (email.lower().strip(),))
            if cur.fetchone():
                conn.close()
                return False, "Este e-mail já está cadastrado."
            cur.execute("INSERT INTO usuarios (nome, email, senha_hash, apikey, tipo, status) VALUES (%s, %s, %s, %s, 'User', 'Ativo')", (nome.strip(), email.lower().strip(), senha, apikey))
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return True, "Conta criada com sucesso! Faça login com seu e-mail e senha."
    except psycopg2.IntegrityError as e:
        if conn: conn.rollback(); conn.close()
        return False, f"Erro de integridade: {str(e)}"
    except Exception as e:
        if conn: conn.close()
        return False, f"Erro ao criar conta: {str(e)}"

def solicitar_reset_senha(email: str) -> Tuple[bool, str, Optional[str]]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado.", None
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT id, nome FROM usuarios WHERE email = %s", (email.lower().strip(),))
            usuario = cur.fetchone()
            if not usuario:
                conn.close()
                return False, "E-mail não cadastrado no sistema.", None
            token = secrets.token_urlsafe(32)
            expires = datetime.now(timezone.utc) + timedelta(hours=1)
            cur.execute("UPDATE usuarios SET reset_token = %s, reset_token_expires = %s WHERE id = %s", (token, expires, usuario['id']))
            conn.commit()
        conn.close()
        return True, f"Token de recuperação gerado para {usuario['nome']}.", token
    except Exception as e:
        if conn: conn.close()
        return False, f"Erro ao processar solicitação: {str(e)}", None

def redefinir_senha(token: str, nova_senha: str) -> Tuple[bool, str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado."
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT id, reset_token, reset_token_expires FROM usuarios WHERE reset_token = %s", (token.strip(),))
            usuario = cur.fetchone()
            if not usuario:
                conn.close()
                return False, "Token inválido ou expirado."
            if usuario['reset_token_expires'] < datetime.now(timezone.utc):
                conn.close()
                return False, "Token expirado. Solicite um novo."
            cur.execute("UPDATE usuarios SET senha_hash = %s, reset_token = NULL, reset_token_expires = NULL WHERE id = %s", (nova_senha, usuario['id']))
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return True, "Senha redefinida com sucesso! Faça login com a nova senha."
    except Exception as e:
        if conn: conn.close()
        return False, f"Erro ao redefinir senha: {str(e)}"

# =============================================================================
# CONSULTAS CACHEADAS AO BANCO NEON
# =============================================================================
def sync_projects_from_azure(pat: str) -> List[str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url or not pat: return []
    try:
        url = f"{ADO_BASE_URL}/_apis/projects?api-version=7.1&$top=100"
        session = get_http_session()
        response = session.get(url, headers=get_ado_headers(pat), timeout=20)
        if response.status_code != 200: return []
        
        azure_projects = response.json().get("value", [])
        project_names = [p["name"] for p in azure_projects]
        
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor() as cur:
            for name in project_names:
                cur.execute("SELECT id FROM devops_projects WHERE name = %s", (name,))
                if not cur.fetchone():
                    cur.execute("INSERT INTO devops_projects (name, is_active) VALUES (%s, TRUE)", (name,))
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return project_names
    except Exception as e:
        logger.error(f"Erro ao sincronizar projetos: {e}")
        return []

@st.cache_data(ttl=300, show_spinner=False)
def get_active_projects_from_db(pat: str = None) -> List[str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT name FROM devops_projects WHERE is_active = TRUE ORDER BY name")
            projects = [row['name'] for row in cur.fetchall()]
        conn.close()
        if not projects and pat:
            projects = sync_projects_from_azure(pat)
        return projects
    except Exception as e:
        logger.error(f"Erro ao buscar projetos: {e}")
        return sync_projects_from_azure(pat) if pat else []

@st.cache_data(ttl=300, show_spinner=False)
def get_colaboradores_from_db() -> pd.DataFrame:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return pd.DataFrame()
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("""
                SELECT id, colaborador_azure, colaborador_ponto, cargo, nivel, 
                       data_inicio, email, promocao, reposicao, data_saida, status
                FROM colaboradores 
                ORDER BY colaborador_ponto
            """)
            rows = cur.fetchall()
            df = pd.DataFrame(rows)
        conn.close()
        return df
    except Exception as e:
        logger.error(f"Erro ao buscar colaboradores: {e}")
        return pd.DataFrame()

@st.cache_data(ttl=120, show_spinner=False)
def get_horasponto_from_db(ano: int, mes: int) -> pd.DataFrame:
    db_url = os.getenv("NEON_DATABASE_URL")
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT id_colaborador, colaborador_ponto, horas_ponto, afastamento FROM horasponto WHERE ano_base = %s AND mes_base = %s", (ano, mes))
            df = pd.DataFrame(cur.fetchall())
        conn.close()
        return df
    except: return pd.DataFrame()

@st.cache_data(ttl=120, show_spinner=False)
def get_horasponto_by_periodo(ano: int, mes: int) -> pd.DataFrame:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return pd.DataFrame()
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT column_name FROM information_schema.columns WHERE table_name = 'horasponto'")
            colunas_existentes = [row['column_name'] for row in cur.fetchall()]
            
            colunas_select = ['id', 'ano_base', 'mes_base', 'colaborador_ponto', 'horas_ponto']
            if 'id_colaborador' in colunas_existentes: colunas_select.append('id_colaborador')
            if 'afastamento' in colunas_existentes: colunas_select.append('afastamento')
            
            cur.execute(f"SELECT {', '.join(colunas_select)} FROM horasponto WHERE ano_base = %s AND mes_base = %s ORDER BY colaborador_ponto", (ano, mes))
            rows = cur.fetchall()
        conn.close()
        if not rows: return pd.DataFrame()
        df = pd.DataFrame(rows)
        if 'afastamento' not in df.columns: df['afastamento'] = 0.0
        
        df_colaboradores = get_colaboradores_from_db()
        if not df_colaboradores.empty:
            colab_map = {str(r['colaborador_ponto']).strip().lower(): {'cargo': r.get('cargo', ''), 'nivel': r.get('nivel', ''), 'id': r.get('id')} for _, r in df_colaboradores.iterrows()}
            df['cargo'] = df['colaborador_ponto'].apply(lambda x: colab_map.get(str(x).strip().lower(), {}).get('cargo', ''))
            df['nivel'] = df['colaborador_ponto'].apply(lambda x: colab_map.get(str(x).strip().lower(), {}).get('nivel', ''))
        return df
    except Exception as e:
        logger.error(f"Erro ao buscar horasponto do período: {e}")
        return pd.DataFrame()

def deletar_horasponto_by_periodo(ano: int, mes: int) -> Tuple[bool, str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado."
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor() as cur:
            cur.execute("DELETE FROM horasponto WHERE ano_base = %s AND mes_base = %s", (ano, mes))
            deletados = cur.rowcount
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return True, f"{deletados} registros deletados com sucesso."
    except Exception as e:
        if conn: conn.rollback(); conn.close()
        return False, f"Erro ao deletar: {str(e)}"

# Inserção em lote otimizada (reduz 85 conexões para apenas 1 conexão)
def inserir_horasponto_lote(registros: List[Dict]) -> Tuple[int, int, List[str]]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url or not registros: return 0, 0, ["Base vazia ou sem URL"]
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=15)
        with conn.cursor() as cur:
            cur.execute("SELECT column_name FROM information_schema.columns WHERE table_name = 'horasponto'")
            colunas_existentes = [r[0] for r in cur.fetchall()]
            tem_afastamento = 'afastamento' in colunas_existentes
            
            cur.execute("SELECT COALESCE(MAX(id), 0) FROM horasponto")
            start_id = cur.fetchone()[0] + 1
            
            dados_insert = []
            for i, r in enumerate(registros):
                row_id = start_id + i
                if tem_afastamento:
                    dados_insert.append((row_id, r['ano'], r['mes'], r['id_colaborador'], r['colaborador_ponto'], r['horas_ponto'], r['afastamento']))
                else:
                    dados_insert.append((row_id, r['ano'], r['mes'], r['id_colaborador'], r['colaborador_ponto'], r['horas_ponto']))
            
            cols = "id, ano_base, mes_base, id_colaborador, colaborador_ponto, horas_ponto" + (", afastamento" if tem_afastamento else "")
            query = f"INSERT INTO horasponto ({cols}) VALUES %s"
            execute_values(cur, query, dados_insert)
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return len(dados_insert), 0, []
    except Exception as e:
        if conn: conn.rollback(); conn.close()
        return 0, len(registros), [str(e)]

@st.cache_data(ttl=300, show_spinner=False)
def get_colaboradores_map() -> Dict[str, int]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return {}
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT id, colaborador_ponto FROM colaboradores")
            rows = cur.fetchall()
        conn.close()
        return {str(row['colaborador_ponto']).strip().lower(): row['id'] for row in rows}
    except Exception as e:
        logger.error(f"Erro ao buscar mapa de colaboradores: {e}")
        return {}

def encontrar_colaborador_id(nome_ponto: str, colaboradores_map: Dict[str, int]) -> Optional[int]:
    if not nome_ponto: return None
    nome_limpo = str(nome_ponto).strip().lower()
    if nome_limpo in colaboradores_map: return colaboradores_map[nome_limpo]
    for nome_db, colab_id in colaboradores_map.items():
        if nome_limpo in nome_db or nome_db in nome_limpo: return colab_id
    palavras_nome = set(nome_limpo.split())
    for nome_db, colab_id in colaboradores_map.items():
        if len(palavras_nome & set(nome_db.split())) >= 2: return colab_id
    return None

@st.cache_data(ttl=600, show_spinner=False)
def get_dias_uteis_from_db(ano: int, mes: int) -> float:
    db_url = os.getenv("NEON_DATABASE_URL")
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT horas_uteis FROM dias_uteis_mensais WHERE ano = %s AND mes = %s LIMIT 1", (ano, mes))
            row = cur.fetchone()
        conn.close()
        return float(row['horas_uteis']) if row else 0.0
    except: return 0.0

@st.cache_data(ttl=600, show_spinner=False)
def get_valores_from_db() -> pd.DataFrame:
    db_url = os.getenv("NEON_DATABASE_URL")
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT nivel, valor FROM valores")
            df = pd.DataFrame(cur.fetchall())
        conn.close()
        return df
    except: return pd.DataFrame(columns=["nivel", "valor"])

@st.cache_data(ttl=600, show_spinner=False)
def get_all_capacities() -> Dict[str, float]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return {}
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT nivel, capacity FROM capacity")
            rows = cur.fetchall()
        conn.close()
        return {str(row['nivel']).strip(): float(row['capacity']) for row in rows}
    except Exception as e:
        logger.error(f"Erro ao buscar capacities: {e}")
        return {}

def limpar_valor_monetario(valor) -> float:
    if valor is None or pd.isna(valor): return 0.0
    valor_str = str(valor).strip().replace('$', '').replace('R$', '').replace(' ', '')
    if ',' in valor_str and '.' in valor_str:
        if valor_str.rfind(',') > valor_str.rfind('.'): valor_str = valor_str.replace('.', '').replace(',', '.')
        else: valor_str = valor_str.replace(',', '')
    elif ',' in valor_str: valor_str = valor_str.replace(',', '.')
    try: return float(valor_str)
    except: return 0.0

@st.cache_data(ttl=60, show_spinner=False)
def listar_usuarios(tipo_usuario: str, usuario_id: int) -> List[Dict]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return []
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            if tipo_usuario == 'Admin': cur.execute("SELECT id, nome, email, tipo, status, apikey, criado_em FROM usuarios ORDER BY nome")
            else: cur.execute("SELECT id, nome, email, tipo, status, apikey, criado_em FROM usuarios WHERE id = %s", (usuario_id,))
            usuarios = cur.fetchall()
        conn.close()
        return usuarios
    except: return []

def atualizar_usuario(usuario_id: int, dados: Dict, tipo_usuario_logado: str, usuario_logado_id: int) -> Tuple[bool, str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado."
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor() as cur:
            if tipo_usuario_logado != 'Admin' and usuario_id != usuario_logado_id:
                conn.close()
                return False, "Você não tem permissão para editar este usuário."
            campos_permitidos = ['nome', 'email', 'senha_hash', 'apikey', 'tipo', 'status'] if tipo_usuario_logado == 'Admin' else ['nome', 'email', 'senha_hash', 'apikey']
            campos_atualizar = {k: v for k, v in dados.items() if k in campos_permitidos and v is not None}
            if not campos_atualizar:
                conn.close()
                return False, "Nenhum campo válido para atualizar."
            set_clause = ", ".join([f"{campo} = %s" for campo in campos_atualizar.keys()])
            valores = list(campos_atualizar.values()) + [usuario_id]
            cur.execute(f"UPDATE usuarios SET {set_clause} WHERE id = %s", valores)
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return True, "Dados atualizados com sucesso!"
    except Exception as e:
        if conn: conn.rollback(); conn.close()
        return False, f"Erro ao atualizar: {str(e)}"

def deletar_usuario(usuario_id: int, tipo_usuario_logado: str) -> Tuple[bool, str]:
    if tipo_usuario_logado != 'Admin': return False, "Apenas administradores podem deletar usuários."
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado."
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor() as cur:
            cur.execute("DELETE FROM usuarios WHERE id = %s", (usuario_id,))
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return True, "Usuário deletado com sucesso!"
    except Exception as e:
        if conn: conn.rollback(); conn.close()
        return False, f"Erro ao deletar: {str(e)}"

# =============================================================================
# CRUD DE COLABORADORES
# =============================================================================
@st.cache_data(ttl=120, show_spinner=False)
def listar_colaboradores() -> pd.DataFrame:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return pd.DataFrame()
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("""
                SELECT id, colaborador_azure, colaborador_ponto, cargo, nivel,
                       data_inicio, email, promocao, reposicao, data_saida, status
                FROM colaboradores ORDER BY colaborador_ponto
            """)
            rows = cur.fetchall()
        conn.close()
        return pd.DataFrame(rows)
    except Exception as e:
        logger.error(f"Erro ao listar colaboradores: {e}")
        return pd.DataFrame()

def criar_colaborador(dados: Dict) -> Tuple[bool, str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado."
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor() as cur:
            cur.execute("SELECT id FROM colaboradores WHERE colaborador_ponto = %s", (dados.get('colaborador_ponto', '').strip(),))
            if cur.fetchone():
                conn.close()
                return False, "Já existe um colaborador cadastrado com este nome."
            cur.execute("""
                INSERT INTO colaboradores 
                (colaborador_azure, colaborador_ponto, cargo, nivel, data_inicio, email, promocao, reposicao, data_saida, status)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, 'Ativo')
            """, (
                dados.get('colaborador_azure', '').strip(),
                dados.get('colaborador_ponto', '').strip(),
                dados.get('cargo', '').strip(),
                dados.get('nivel', '').strip(),
                dados.get('data_inicio') if dados.get('data_inicio') else None,
                dados.get('email', '').strip(),
                dados.get('promocao') if dados.get('promocao') else None,
                dados.get('reposicao', '').strip(),
                dados.get('data_saida') if dados.get('data_saida') else None
            ))
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return True, "Colaborador cadastrado com sucesso!"
    except Exception as e:
        if conn: conn.rollback(); conn.close()
        return False, f"Erro ao cadastrar: {str(e)}"

def atualizar_colaborador(colaborador_id: int, dados: Dict) -> Tuple[bool, str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado."
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor() as cur:
            cur.execute("""
                UPDATE colaboradores SET
                    colaborador_azure = %s, colaborador_ponto = %s, cargo = %s, nivel = %s,
                    data_inicio = %s, email = %s, promocao = %s, reposicao = %s, data_saida = %s, status = %s
                WHERE id = %s
            """, (
                dados.get('colaborador_azure', '').strip(),
                dados.get('colaborador_ponto', '').strip(),
                dados.get('cargo', '').strip(),
                dados.get('nivel', '').strip(),
                dados.get('data_inicio') if dados.get('data_inicio') else None,
                dados.get('email', '').strip(),
                dados.get('promocao') if dados.get('promocao') else None,
                dados.get('reposicao', '').strip(),
                dados.get('data_saida') if dados.get('data_saida') else None,
                dados.get('status', 'Ativo'),
                colaborador_id
            ))
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return True, "Colaborador atualizado com sucesso!"
    except Exception as e:
        if conn: conn.rollback(); conn.close()
        return False, f"Erro ao atualizar: {str(e)}"

def deletar_colaborador(colaborador_id: int) -> Tuple[bool, str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado."
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor() as cur:
            cur.execute("DELETE FROM colaboradores WHERE id = %s", (colaborador_id,))
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return True, "Colaborador deletado com sucesso!"
    except Exception as e:
        if conn: conn.rollback(); conn.close()
        return False, f"Erro ao deletar: {str(e)}"

# =============================================================================
# CRUD DE PROJETOS
# =============================================================================
@st.cache_data(ttl=120, show_spinner=False)
def listar_projetos() -> pd.DataFrame:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return pd.DataFrame()
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT column_name FROM information_schema.columns WHERE table_name = 'devops_projects'")
            colunas_existentes = [row['column_name'] for row in cur.fetchall()]
            
            colunas_select = ['id', 'name', 'is_active']
            if 'created_at' in colunas_existentes: colunas_select.append('created_at')
            elif 'criado_em' in colunas_existentes: colunas_select.append('criado_em')
            if 'updated_at' in colunas_existentes: colunas_select.append('updated_at')
            elif 'atualizado_em' in colunas_existentes: colunas_select.append('atualizado_em')
            
            cur.execute(f"SELECT {', '.join(colunas_select)} FROM devops_projects ORDER BY name")
            rows = cur.fetchall()
        conn.close()
        return pd.DataFrame(rows)
    except Exception as e:
        logger.error(f"Erro ao listar projetos: {e}")
        return pd.DataFrame()

def criar_projeto(name: str, is_active: bool = True) -> Tuple[bool, str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado."
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor() as cur:
            cur.execute("SELECT id FROM devops_projects WHERE name = %s", (name.strip(),))
            if cur.fetchone():
                conn.close()
                return False, "Já existe um projeto cadastrado com este nome."
            cur.execute("INSERT INTO devops_projects (name, is_active) VALUES (%s, %s)", (name.strip(), is_active))
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return True, "Projeto cadastrado com sucesso!"
    except Exception as e:
        if conn: conn.rollback(); conn.close()
        return False, f"Erro ao cadastrar: {str(e)}"

def atualizar_projeto(projeto_id: int, name: str, is_active: bool) -> Tuple[bool, str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado."
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor() as cur:
            cur.execute("SELECT column_name FROM information_schema.columns WHERE table_name = 'devops_projects' AND column_name IN ('updated_at', 'atualizado_em')")
            col_timestamp = cur.fetchone()
            if col_timestamp:
                cur.execute(f"UPDATE devops_projects SET name = %s, is_active = %s, {col_timestamp[0]} = CURRENT_TIMESTAMP WHERE id = %s", (name.strip(), is_active, projeto_id))
            else:
                cur.execute("UPDATE devops_projects SET name = %s, is_active = %s WHERE id = %s", (name.strip(), is_active, projeto_id))
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return True, "Projeto atualizado com sucesso!"
    except Exception as e:
        if conn: conn.rollback(); conn.close()
        return False, f"Erro ao atualizar: {str(e)}"

def deletar_projeto(projeto_id: int) -> Tuple[bool, str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado."
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor() as cur:
            cur.execute("DELETE FROM devops_projects WHERE id = %s", (projeto_id,))
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return True, "Projeto deletado com sucesso!"
    except Exception as e:
        if conn: conn.rollback(); conn.close()
        return False, f"Erro ao deletar: {str(e)}"

# =============================================================================
# AZURE DEVOPS OTIMIZADO COM CACHE E PARALELIZAÇÃO
# =============================================================================
def get_ado_headers(pat: str) -> Dict[str, str]:
    auth_string = f":{pat}"
    return {"Authorization": f"Basic {base64.b64encode(auth_string.encode('ascii')).decode('ascii')}", "Content-Type": "application/json"}

@st.cache_data(ttl=600, show_spinner=False)
def validate_ado_api_key(pat: str) -> Tuple[bool, str]:
    if not pat: return False, "API Key vazia."
    url = f"{ADO_BASE_URL}/_apis/projects?api-version=7.1&$top=1"
    try:
        session = get_http_session()
        response = session.get(url, headers=get_ado_headers(pat), timeout=15)
        if response.status_code == 401: return False, "API Azure inválida: autenticação não autorizada (401)."
        elif response.status_code == 403: return False, "API Azure inválida: acesso proibido (403)."
        elif response.status_code == 404: return False, "API Azure inválida: organização não encontrada (404)."
        elif response.status_code != 200: return False, f"API Azure inválida: erro {response.status_code}."
        return True, "API Key válida."
    except Exception as e:
        return False, f"API Azure inválida: {str(e)}"

def parse_date(date_str: str) -> Optional[datetime]:
    if not date_str: return None
    try: return datetime.fromisoformat(date_str.replace('Z', ''))
    except: return None

@st.cache_data(ttl=300, show_spinner=False)
def fetch_all_sprints_for_project(pat: str, project: str, year: int, month: int) -> List[Dict[str, Any]]:
    url = f"{ADO_BASE_URL}/{urllib.parse.quote(project, safe='')}/_apis/work/teamsettings/iterations?api-version=7.1"
    try:
        session = get_http_session()
        response = session.get(url, headers=get_ado_headers(pat), timeout=15)
        if response.status_code in [401, 403]: return []
        response.raise_for_status()
        data = response.json()
        month_start = datetime(year, month, 1, 0, 0, 0)
        month_end = datetime(year + 1, 1, 1, 0, 0, 0) - timedelta(seconds=1) if month == 12 else datetime(year, month + 1, 1, 0, 0, 0) - timedelta(seconds=1)
        valid_sprints = []
        for item in data.get("value", []):
            attrs = item.get("attributes", {})
            start_date = parse_date(attrs.get("startDate"))
            if start_date and month_start <= start_date <= month_end:
                valid_sprints.append({"project": project, "name": item["name"], "path": item["path"], "start": start_date.strftime("%d/%m/%Y"), "finish": (attrs.get("finishDate") or "").replace("T", " ").split(".")[0]})
        return valid_sprints
    except: return []

@st.cache_data(ttl=300, show_spinner=False)
def fetch_all_workitems_ids_by_project(pat: str, project: str, sprint_paths: Tuple[str, ...]) -> Tuple[List[int], Optional[str]]:
    if not sprint_paths: return [], "Nenhum path"
    formatted_paths = ", ".join([f"'{p.replace(chr(92), chr(92)+chr(92))}'" for p in sprint_paths])
    wiql_query = f"SELECT [System.Id] FROM WorkItems WHERE [System.IterationPath] IN ({formatted_paths}) ORDER BY [System.Id]"
    try:
        session = get_http_session()
        response = session.post(f"{ADO_BASE_URL}/{urllib.parse.quote(project, safe='')}/_apis/wit/wiql?api-version=7.1", headers=get_ado_headers(pat), json={"query": wiql_query}, timeout=30)
        if response.status_code in [401, 403]: return [], "API Key inválida"
        response.raise_for_status()
        return [w["id"] for w in response.json().get("workItems", [])], None
    except: return [], "Erro"

@st.cache_data(ttl=300, show_spinner=False)
def fetch_hierarchy_links_by_project(pat: str, project: str, sprint_paths: Tuple[str, ...]) -> Tuple[List[Dict], Optional[str]]:
    if not sprint_paths: return [], "Nenhum path"
    formatted_paths = ", ".join([f"'{p.replace(chr(92), chr(92)+chr(92))}'" for p in sprint_paths])
    wiql_query = f"SELECT [System.Id] FROM workitemLinks WHERE ([Source].[System.IterationPath] IN ({formatted_paths})) AND ([System.Links.LinkType] = 'System.LinkTypes.Hierarchy-Forward' OR [System.Links.LinkType] = 'System.LinkTypes.Hierarchy-Reverse') AND ([Target].[System.IterationPath] IN ({formatted_paths})) MODE (MustContain)"
    try:
        session = get_http_session()
        response = session.post(f"{ADO_BASE_URL}/{urllib.parse.quote(project, safe='')}/_apis/wit/wiql?api-version=7.1", headers=get_ado_headers(pat), json={"query": wiql_query}, timeout=30)
        if response.status_code in [401, 403]: return [], "API Key inválida"
        response.raise_for_status()
        return response.json().get("workItemRelations", []), None
    except: return [], "Erro"

@st.cache_data(ttl=300, show_spinner=False)
def fetch_workitems_details_batch(pat: str, project: str, ids: Tuple[int, ...]) -> Tuple[Dict[int, Dict], Optional[str]]:
    if not ids: return {}, "Nenhum ID"
    url = f"{ADO_BASE_URL}/{urllib.parse.quote(project, safe='')}/_apis/wit/workitemsbatch?api-version=7.1"
    fields = ["System.Id", "System.Title", "System.WorkItemType", "System.State", "System.IterationPath", 
              "System.CreatedDate", "Microsoft.VSTS.Common.ClosedDate", 
              "Microsoft.VSTS.Scheduling.OriginalEstimate", "System.TeamProject", 
              "Microsoft.VSTS.Scheduling.CompletedWork", "System.AssignedTo",
              "Microsoft.VSTS.Common.ActivatedDate"]
    all_items = {}
    try:
        session = get_http_session()
        for i in range(0, len(ids), 200):
            response = session.post(url, headers=get_ado_headers(pat), json={"ids": list(ids[i:i+200]), "fields": fields}, timeout=20)
            if response.status_code in [401, 403]: return all_items, "API Key inválida"
            response.raise_for_status()
            for item in response.json().get("value", []): all_items[item["id"]] = item
        return all_items, None
    except: return all_items, "Erro"

# =============================================================================
# FUNÇÕES AUXILIARES
# =============================================================================
def format_date(date_str: Optional[str]) -> str:
    if not date_str: return "-"
    try: return datetime.fromisoformat(date_str.replace('Z', '+00:00')).strftime("%d/%m/%Y %H:%M:%S")
    except: return "-"

def format_hours(hours: Optional[float]) -> str: return f"{float(hours):.2f}" if hours is not None else "0.00"

def format_moeda_brl(valor: float) -> str:
    return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

def esc(text: str) -> str:
    if text is None: return ""
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;").replace("'", "&#39;")

def state_to_css_class(state: str) -> str: return state.replace(" ", "-").replace("'", "")

def get_assigned_to_name(assigned_to: Any) -> str:
    if not assigned_to: return "Não atribuído"
    if isinstance(assigned_to, dict): return assigned_to.get("displayName", "Não atribuído")
    return str(assigned_to)

def encontrar_colaborador_ponto(assigned_name, azure_to_ponto):
    if assigned_name == "Não atribuído": return assigned_name
    azure_key = assigned_name.strip().lower()
    if "<" in azure_key: azure_key = azure_key.split("<")[0].strip()
    if azure_key in azure_to_ponto: return azure_to_ponto[azure_key]
    for azure_name, ponto_name in azure_to_ponto.items():
        if azure_key in azure_name or azure_name in azure_key: return ponto_name
    palavras_azure = azure_key.split()
    for azure_name, ponto_name in azure_to_ponto.items():
        if len(set(palavras_azure) & set(azure_name.split())) >= 2: return ponto_name
    return assigned_name

# =============================================================================
# EXPORTAÇÃO
# =============================================================================
def generate_csv_parents_only(parents_data, children_map, all_workitems):
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_ALL)
    writer.writerow(["Projeto", "Sprint", "ID", "Título", "Tipo", "Situação", "Horas Previstas", "Link", "Dt Criação", "Dt Término", "Qtd Filhos"])
    for p in parents_data:
        f = p.get("fields", {})
        pid, proj = p["id"], p.get("_project", "")
        sprint, title = str(f.get("System.IterationPath", "N/A")), f.get("System.Title", "")
        ptype, pstate = f.get("System.WorkItemType", ""), f.get("System.State", "")
        hours = format_hours(f.get("Microsoft.VSTS.Scheduling.OriginalEstimate"))
        link = f"{ADO_BASE_URL}/{urllib.parse.quote(proj, safe='')}/_workitems/edit/{pid}"
        writer.writerow([proj, sprint, pid, title, ptype, pstate, hours, link, format_date(f.get("System.CreatedDate")), format_date(f.get("Microsoft.VSTS.Common.ClosedDate")), len(children_map.get(pid, []))])
    return "\ufeff" + output.getvalue()

def generate_csv_complete(parents_data, children_map, all_workitems):
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_ALL)
    writer.writerow(["Projeto", "Sprint", "ID", "Título", "Tipo", "Situação", "Horas Previstas", "Link", "Dt Criação", "Dt Término", "Nível", "ID Pai"])
    for p in parents_data:
        f = p.get("fields", {})
        pid, proj = p["id"], p.get("_project", "")
        sprint, title = str(f.get("System.IterationPath", "N/A")), f.get("System.Title", "")
        ptype, pstate = f.get("System.WorkItemType", ""), f.get("System.State", "")
        hours = format_hours(f.get("Microsoft.VSTS.Scheduling.OriginalEstimate"))
        link = f"{ADO_BASE_URL}/{urllib.parse.quote(proj, safe='')}/_workitems/edit/{pid}"
        writer.writerow([proj, sprint, pid, title, ptype, pstate, hours, link, format_date(f.get("System.CreatedDate")), format_date(f.get("Microsoft.VSTS.Common.ClosedDate")), "PAI", ""])
        for cid in children_map.get(pid, []):
            if cid not in all_workitems: continue
            cf = all_workitems[cid].get("fields", {})
            writer.writerow([proj, sprint, cid, cf.get("System.Title", ""), cf.get("System.WorkItemType", ""), cf.get("System.State", ""), format_hours(cf.get("Microsoft.VSTS.Scheduling.OriginalEstimate")), f"{ADO_BASE_URL}/{urllib.parse.quote(proj, safe='')}/_workitems/edit/{cid}", format_date(cf.get("System.CreatedDate")), format_date(cf.get("Microsoft.VSTS.Common.ClosedDate")), "FILHO", pid])
    return "\ufeff" + output.getvalue()

def generate_csv_backlogs_nao_done(parents_data):
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_ALL)
    writer.writerow(["Projeto", "Sprint", "ID", "Descrição da História", "Assign To", "Status"])
    count = 0
    for p in parents_data:
        f = p.get("fields", {})
        if f.get("System.WorkItemType") == "Product Backlog Item" and f.get("System.State") != "Done":
            proj = p.get("_project", "")
            sprint = str(f.get("System.IterationPath", "N/A")).split("\\")[-1] if f.get("System.IterationPath") else "N/A"
            writer.writerow([proj, sprint, p["id"], f.get("System.Title", ""), get_assigned_to_name(f.get("System.AssignedTo")), f.get("System.State", "")])
            count += 1
    return "\ufeff" + output.getvalue(), count

def generate_csv_tasks_done_sem_horas(parents_data):
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_ALL)
    writer.writerow(["Projeto", "Sprint", "ID", "Descrição da Tarefa", "Assign To", "Status", "Completed Work"])
    count = 0
    for p in parents_data:
        f = p.get("fields", {})
        if f.get("System.WorkItemType") == "Task" and f.get("System.State") == "Done":
            completed_work = f.get("Microsoft.VSTS.Scheduling.CompletedWork")
            completed_work_float = float(completed_work) if completed_work is not None else 0.0
            if completed_work_float == 0:
                proj = p.get("_project", "")
                sprint = str(f.get("System.IterationPath", "N/A")).split("\\")[-1] if f.get("System.IterationPath") else "N/A"
                writer.writerow([proj, sprint, p["id"], f.get("System.Title", ""), get_assigned_to_name(f.get("System.AssignedTo")), f.get("System.State", ""), f"{completed_work_float:.2f}"])
                count += 1
    return "\ufeff" + output.getvalue(), count

def generate_excel_faturamento(df_faturamento: pd.DataFrame, mes: str, ano: int, dias_uteis: float, total_faturado: float) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Faturamento"
    header_font = Font(name='Montserrat', size=14, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='FF6600', end_color='FF6600', fill_type='solid')
    info_font = Font(name='Montserrat', size=11, bold=True, color='333333')
    data_font = Font(name='Montserrat', size=10)
    money_font = Font(name='Montserrat', size=10, bold=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    ws.merge_cells('A1:M1')
    cell_title = ws['A1']
    cell_title.value = f"FATURAMENTO - {mes}/{ano}"
    cell_title.font = header_font
    cell_title.fill = header_fill
    cell_title.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 35
    
    ws.merge_cells('A2:M2')
    cell_info = ws['A2']
    cell_info.value = f"Dias Úteis: {dias_uteis/8:.0f} | Total Faturado: {format_moeda_brl(total_faturado)}"
    cell_info.font = info_font
    cell_info.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 25
    
    headers = ["Profissional", "Cargo", "Nível", "Data de Início", "Promoção de Cargo", "Novo Colaborador", "Data de Saída", "Horas Trabalhadas", "Horas Afast Acordado", "Horas Afast Legal", "Horas Arredondadas", "Valor Profissional", "Valor Recalculado"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header
        cell.font = Font(name='Montserrat', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='2C5282', end_color='2C5282', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[4].height = 35
    
    column_widths = [35, 20, 15, 15, 20, 25, 15, 18, 22, 18, 20, 20, 20]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    
    for row_num, (_, row) in enumerate(df_faturamento.iterrows(), 5):
        for col_idx, col_name in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=row[col_name])
            if col_name in ["Valor Profissional", "Valor Recalculado"]:
                cell.font = money_font
                cell.number_format = 'R$ #,##0.00'
            else:
                cell.font = data_font
            cell.border = thin_border
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

# =============================================================================
# RENDERIZAÇÃO - TABELA HTML
# =============================================================================
def render_table_html(parents_data, children_map, all_workitems, total_hours):
    css = """
    <style>
    .ado-tbl { width: 100%; border-collapse: collapse; font-family: 'Segoe UI', sans-serif; font-size: 13px; background: white; }
    .ado-tbl th { background-color: #FF6600; color: white; padding: 10px 8px; text-align: left; border: 1px solid #ddd; font-size: 12px; }
    .ado-tbl td { padding: 0; border: 1px solid #f0f0f0; }
    .parent-tbl { width: 100%; border-collapse: collapse; background: white; cursor: pointer; }
    .parent-tbl:hover { background: #fff5eb; }
    .parent-tbl td { padding: 10px 8px; border: none; vertical-align: middle; }
    .child-tbl { width: 100%; font-size: 12px; border-collapse: collapse; background: #fafafa; margin-top: 5px; }
    .child-tbl th { background: #e9ecef; color: #333; padding: 6px 8px; text-align: left; font-size: 11px; }
    .child-tbl td { padding: 6px 8px; border-bottom: 1px solid #e0e0e0; border-left: 3px solid #FF6600; }
    .ftr { background: #E0E0E0; font-weight: bold; }
    .ftr td { padding: 12px 8px; border-top: 2px solid #bbb; }
    .lnk { display: inline-block; padding: 4px 10px; background: #0078d4; color: white; text-decoration: none; border-radius: 4px; font-size: 11px; }
    .lnk:hover { background: #005a9e; }
    .badge { display: inline-block; padding: 2px 8px; border-radius: 10px; font-size: 11px; font-weight: 600; }
    .b-Done { background: #d4edda; color: #155724; } .b-Resolved { background: #d1ecf1; color: #0c5460; }
    .b-In-Progress { background: #fff3cd; color: #856404; } .b-To-Do { background: #e2e3e5; color: #383d41; }
    details > summary { list-style: none; cursor: pointer; padding: 0; }
    details > summary::-webkit-details-marker { display: none; }
    details > summary::before { content: '▶'; display: inline-block; color: #FF6600; font-size: 10px; margin-right: 4px; }
    details[open] > summary::before { content: '▼'; }
    .dt { font-size: 11px; color: #666; } .sprint-path { font-size: 11px; color: #444; word-break: break-word; }
    </style>
    """
    rows = []
    for p in parents_data:
        f = p.get("fields", {})
        pid, title = p["id"], esc(f.get("System.Title", ""))
        ptype, pstate = esc(f.get("System.WorkItemType", "")), esc(f.get("System.State", ""))
        hours = format_hours(f.get("Microsoft.VSTS.Scheduling.OriginalEstimate"))
        sprint_path = esc(str(f.get("System.IterationPath", "N/A")))
        created, closed = format_date(f.get("System.CreatedDate")), format_date(f.get("Microsoft.VSTS.Common.ClosedDate"))
        proj = p.get("_project", "")
        link = f"{ADO_BASE_URL}/{urllib.parse.quote(proj, safe='')}/_workitems/edit/{pid}"
        child_ids = children_map.get(pid, [])
        children = [all_workitems[cid] for cid in child_ids if cid in all_workitems]
        state_class = state_to_css_class(pstate)
        
        parent_row = f"""<tr><td colspan="10" style="padding:0;border:none;">
<details><summary><table class="parent-tbl"><tr>
<td style="width:3%"></td><td style="width:22%" class="sprint-path"><strong>{sprint_path}</strong></td>
<td style="width:6%;color:#FF6600;font-weight:700;">#{pid}</td><td style="width:25%" title="{title}">{title[:50]}{'...' if len(title)>50 else ''}</td>
<td style="width:10%">{ptype}</td><td style="width:10%"><span class="badge b-{state_class}">{pstate}</span></td>
<td style="width:7%;font-weight:600;color:#2c5282;">{hours}</td>
<td style="width:6%"><a href="{link}" target="_blank" class="lnk" onclick="event.stopPropagation()">Abrir</a></td>
<td style="width:6%" class="dt">{created}</td><td style="width:6%" class="dt">{closed}</td>
</tr></table></summary>"""
        if children:
            parent_row += '<div style="padding:0 0 0 30px;"><table class="child-tbl"><thead><tr><th style="width:3%"></th><th style="width:22%">Sprint</th><th style="width:6%">ID</th><th style="width:25%">Título</th><th style="width:10%">Tipo</th><th style="width:10%">Situação</th><th style="width:7%">Horas</th><th style="width:6%">Link</th><th style="width:6%">Dt Criação</th><th style="width:6%">Dt Término</th></tr></thead><tbody>'
            for c in children:
                cf = c.get("fields", {})
                cstate = esc(cf.get("System.State", ""))
                parent_row += f"""<tr><td></td><td class="sprint-path">{sprint_path}</td><td style="color:#666;">#{c['id']}</td>
                <td>{esc(cf.get('System.Title',''))}</td><td>{esc(cf.get('System.WorkItemType',''))}</td>
                <td><span class="badge b-{state_to_css_class(cstate)}">{cstate}</span></td>
                <td>{format_hours(cf.get('Microsoft.VSTS.Scheduling.OriginalEstimate'))}</td><td>-</td>
                <td class="dt">{format_date(cf.get('System.CreatedDate'))}</td><td class="dt">{format_date(cf.get('Microsoft.VSTS.Common.ClosedDate'))}</td></tr>"""
            parent_row += "</tbody></table></div>"
        else:
            parent_row += '<div style="padding:10px 0 10px 30px;color:#888;font-style:italic;font-size:12px;">Nenhuma atividade filha vinculada.</div>'
        parent_row += "</details></td></tr>"
        rows.append(parent_row)
    
    full_html = css + f"""<table class="ado-tbl"><thead><tr>
<th style="width:3%"></th><th style="width:22%">Sprint</th><th style="width:6%">ID</th><th style="width:25%">Título</th>
<th style="width:10%">Tipo</th><th style="width:10%">Situação</th><th style="width:7%">Horas</th><th style="width:6%">Link</th>
<th style="width:6%">Dt Criação</th><th style="width:6%">Dt Término</th></tr></thead><tbody>\n""" + "\n".join(rows) + f"""
<tr class="ftr"><td colspan="6" style="text-align:right;">TOTAL ACUMULADO DE HORAS PREVISTAS:</td>
<td style="color:#FF6600;font-size:16px;">{total_hours:.2f} h</td><td colspan="3"></td></tr></tbody></table>"""
    st.markdown(full_html, unsafe_allow_html=True)

# =============================================================================
# RENDERIZAÇÃO - ABAS
# =============================================================================
def render_colaborador_tab(all_workitems_colab, project_of_item, df_colaboradores, year, month_name):
    periodo_atual = f"{year}-{month_name}"
    if st.session_state.get('filtro_colab_periodo') != periodo_atual:
        st.session_state.filtro_colab_periodo = periodo_atual
        if 'filtro_colab_tab' in st.session_state: del st.session_state.filtro_colab_tab
    
    azure_to_ponto = {}
    for _, row in df_colaboradores.iterrows():
        azure_name = str(row["colaborador_azure"]).strip().lower()
        ponto_name = str(row["colaborador_ponto"]).strip()
        if azure_name: azure_to_ponto[azure_name] = ponto_name
    
    tarefas_list = []
    for item_id, item in all_workitems_colab.items():
        f = item.get("fields", {})
        if f.get("System.State") != "Done": continue
        assigned_to = f.get("System.AssignedTo")
        assigned_name = get_assigned_to_name(assigned_to)
        colaborador_ponto = encontrar_colaborador_ponto(assigned_name, azure_to_ponto)
        proj = project_of_item.get(item_id, "")
        completed_work = float(f.get("Microsoft.VSTS.Scheduling.CompletedWork") or 0)
        link = f"{ADO_BASE_URL}/{urllib.parse.quote(proj, safe='')}/_workitems/edit/{item_id}"
        tarefas_list.append({"Colaborador": colaborador_ponto, "Projeto": proj, "Sprint": str(f.get("System.IterationPath", "N/A")).split("\\")[-1],
            "ID": item_id, "Tipo": f.get("System.WorkItemType", ""), "Título": f.get("System.Title", ""), "Horas Apontadas": completed_work, "Link": link})
    
    if not tarefas_list:
        st.warning("⚠️ Nenhuma tarefa com State=Done encontrada.")
        return
    
    df_tarefas = pd.DataFrame(tarefas_list)
    filtro = st.sidebar.multiselect("👥 Filtrar por Colaborador", options=sorted(df_tarefas["Colaborador"].unique()), default=sorted(df_tarefas["Colaborador"].unique()), key="filtro_colab_tab")
    if filtro: df_tarefas = df_tarefas[df_tarefas["Colaborador"].isin(filtro)]
    
    df_sumario = df_tarefas.groupby("Colaborador").agg({"Horas Apontadas": "sum", "ID": "count"}).reset_index()
    df_sumario.columns = ["Colaborador", "Total Horas", "Qtd Tarefas"]
    df_sumario["Média Horas/Tarefa"] = (df_sumario["Total Horas"] / df_sumario["Qtd Tarefas"]).round(2)
    df_sumario = df_sumario.sort_values("Colaborador", ascending=True)
    
    st.subheader("📊 Sumário por Colaborador (State: Done)")
    st.info(f"Total: **{len(df_tarefas)} tarefa(s)** | **{df_tarefas['Horas Apontadas'].sum():.2f}h**")
    st.dataframe(df_sumario, width="stretch", hide_index=True,
        column_config={
            "Colaborador": st.column_config.TextColumn("Colaborador", width="large"),
            "Total Horas": st.column_config.NumberColumn("Total Horas", format="%.2f h"),
            "Qtd Tarefas": st.column_config.NumberColumn("Qtd Tarefas"),
            "Média Horas/Tarefa": st.column_config.NumberColumn("Média Horas/Tarefa", format="%.2f h", help="Média de horas apontadas por tarefa concluída")
        })
    
    st.markdown("---")
    st.subheader("📈 Horas Apontadas por Colaborador")
    fig = px.bar(df_sumario, x="Colaborador", y="Total Horas", text="Total Horas", color="Colaborador", color_discrete_sequence=px.colors.qualitative.Set3)
    fig.update_layout(margin=dict(t=20, b=80, l=20, r=20), height=400, showlegend=False, xaxis_tickangle=-45)
    fig.update_traces(texttemplate='%{text:.1f}h', textposition='outside')
    st.plotly_chart(fig, width="stretch")
    
    st.markdown("---")
    st.subheader("📋 Detalhamento por Colaborador")
    for colaborador in sorted(df_tarefas["Colaborador"].unique()):
        df_col = df_tarefas[df_tarefas["Colaborador"] == colaborador].sort_values(["Projeto", "Sprint"])
        with st.expander(f"👤 **{colaborador}** — {len(df_col)} tarefa(s) | {df_col['Horas Apontadas'].sum():.2f}h", expanded=False):
            st.dataframe(df_col[["Projeto", "Sprint", "ID", "Tipo", "Título", "Horas Apontadas", "Link"]], width="stretch", hide_index=True,
                column_config={"Projeto": st.column_config.TextColumn("Projeto"), "Sprint": st.column_config.TextColumn("Sprint"), "ID": st.column_config.NumberColumn("ID"),
                    "Tipo": st.column_config.TextColumn("Tipo"), "Título": st.column_config.TextColumn("Título", width="large"),
                    "Horas Apontadas": st.column_config.NumberColumn("Horas", format="%.2f h"), "Link": st.column_config.LinkColumn("Link", display_text="Abrir")})

def render_horas_colaborador_tab(df_colaboradores, df_horasponto, horas_uteis_mensais, all_workitems_colab, project_of_item, year, month_name):
    periodo_atual = f"{year}-{month_name}"
    if st.session_state.get('filtro_horas_periodo') != periodo_atual:
        st.session_state.filtro_horas_periodo = periodo_atual
        if 'filtro_horas' in st.session_state: del st.session_state.filtro_horas
    
    colaborador_nivel = {}
    for _, row in df_colaboradores.iterrows():
        ponto_name = str(row["colaborador_ponto"]).strip()
        nivel = str(row.get("nivel", "")).strip()
        if ponto_name: colaborador_nivel[ponto_name] = nivel
    
    # Busca todas as capacities de uma única vez em cache
    nivel_capacity = get_all_capacities()
    
    horas_azure = {}
    azure_to_ponto = {}
    for _, row in df_colaboradores.iterrows():
        azure_name = str(row["colaborador_azure"]).strip().lower()
        ponto_name = str(row["colaborador_ponto"]).strip()
        if azure_name: azure_to_ponto[azure_name] = ponto_name
    
    for item_id, item in all_workitems_colab.items():
        f = item.get("fields", {})
        if f.get("System.State") != "Done": continue
        assigned_name = get_assigned_to_name(f.get("System.AssignedTo"))
        colaborador_ponto = encontrar_colaborador_ponto(assigned_name, azure_to_ponto)
        horas_azure[colaborador_ponto] = horas_azure.get(colaborador_ponto, 0) + float(f.get("Microsoft.VSTS.Scheduling.CompletedWork") or 0)
    
    lista_consolidada = []
    for _, hp_row in df_horasponto.iterrows():
        ponto_name = str(hp_row["colaborador_ponto"]).strip()
        horas_ponto = float(hp_row["horas_ponto"]) if pd.notna(hp_row.get("horas_ponto")) else 0.0
        afastamento = float(hp_row["afastamento"]) if pd.notna(hp_row.get("afastamento")) else 0.0
        nivel = colaborador_nivel.get(ponto_name, "")
        capacity = nivel_capacity.get(nivel, 0.0)
        capacity_calculado = (horas_ponto / 8.0) * capacity if (horas_ponto > 0 and capacity > 0) else 0.0
        
        horas_azure_colab = horas_azure.get(ponto_name, 0.0)
        diferenca = horas_azure_colab - horas_ponto
        
        lista_consolidada.append({
            "Colaborador": ponto_name, 
            "Ponto": horas_ponto, 
            "Azure": horas_azure_colab, 
            "Afastamento": afastamento, 
            "Capacity": capacity_calculado,
            "Diferença": round(diferenca, 2)
        })
    
    df_consolidado = pd.DataFrame(lista_consolidada).sort_values("Colaborador", ascending=True)
    filtro = st.sidebar.multiselect("👥 Filtrar por Colaborador", options=sorted(df_consolidado["Colaborador"].unique()), default=sorted(df_consolidado["Colaborador"].unique()), key="filtro_horas")
    if filtro: df_consolidado = df_consolidado[df_consolidado["Colaborador"].isin(filtro)]
    
    col1, col2, col3, col4, col5 = st.columns(5)
    col1.metric("🕐 Total Ponto", f"{df_consolidado['Ponto'].sum():.2f} h")
    col2.metric("💻 Total Azure", f"{df_consolidado['Azure'].sum():.2f} h")
    col3.metric("🏖️ Total Afastamento", f"{df_consolidado['Afastamento'].sum():.2f} h")
    col4.metric("📊 Total Capacity", f"{df_consolidado['Capacity'].sum():.2f} h")
    
    diff_total = df_consolidado['Diferença'].sum()
    col5.metric("⚖️ Diferença Total", f"{diff_total:+.2f} h", 
                delta=f"{diff_total:+.2f}h",
                delta_color="normal" if abs(diff_total) < 10 else ("inverse" if diff_total < 0 else "normal"))
    
    st.markdown("---")
    st.subheader("🕐 Horas por Colaborador")
    st.info(f"Período: **{month_name}/{year}** | Horas úteis: **{horas_uteis_mensais:.2f}h**")
    
    def colorir_diferenca(val):
        if pd.isna(val): return ''
        if val > 0: return 'background-color: #d4edda; color: #155724'
        elif val < 0: return 'background-color: #f8d7da; color: #721c24'
        else: return 'background-color: #fff3cd; color: #856404'
    
    styler = df_consolidado.style.map(colorir_diferenca, subset=['Diferença'])
    
    st.dataframe(
        styler,
        width="stretch", 
        hide_index=True,
        column_config={
            "Colaborador": st.column_config.TextColumn("Colaborador", width="large"), 
            "Ponto": st.column_config.NumberColumn("Ponto (h)", format="%.2f h"),
            "Azure": st.column_config.NumberColumn("Azure (h)", format="%.2f h"), 
            "Afastamento": st.column_config.NumberColumn("Afastamento (h)", format="%.2f h"),
            "Capacity": st.column_config.NumberColumn("Capacity (h)", format="%.2f h"),
            "Diferença": st.column_config.NumberColumn("Diferença (h)", format="%+.2f h", help="Positivo = excesso no Azure | Negativo = falta no Azure")
        }
    )
    
    st.markdown("---")
    st.subheader("📈 Comparativo: Ponto vs Azure vs Capacity")
    df_grafico = df_consolidado.melt(id_vars=["Colaborador"], value_vars=["Ponto", "Azure", "Capacity"], var_name="Fonte", value_name="Horas")
    fig = px.bar(df_grafico, x="Colaborador", y="Horas", color="Fonte", barmode="group", color_discrete_map={"Ponto": "#28a745", "Azure": "#0078d4", "Capacity": "#FF6600"})
    fig.update_layout(margin=dict(t=20, b=80, l=20, r=20), height=450, xaxis_tickangle=-45)
    st.plotly_chart(fig, width="stretch")
    
    st.markdown("---")
    st.subheader("⚖️ Diferença entre Azure e Ponto por Colaborador")
    st.info("🟢 **Verde**: Horas excedentes no Azure | 🔴 **Vermelho**: Horas faltantes no Azure")
    
    df_diff = df_consolidado.copy().sort_values('Diferença', ascending=False)
    fig_diff = px.bar(df_diff, x="Colaborador", y="Diferença", color="Diferença", color_continuous_scale=['#dc3545', '#ffc107', '#28a745'], text="Diferença")
    fig_diff.update_layout(margin=dict(t=20, b=80, l=20, r=20), height=450, xaxis_tickangle=-45, coloraxis_showscale=False)
    fig_diff.update_traces(texttemplate='%{text:+.1f}h', textposition='outside')
    fig_diff.add_hline(y=0, line_dash="solid", line_color="black", line_width=2)
    st.plotly_chart(fig_diff, width="stretch")

# =============================================================================
# PONTO ELETRÔNICO
# =============================================================================
def render_ponto_tab(year, month_name, month_num, tipo_usuario):
    st.subheader("🕐 Gerenciamento de Ponto Eletrônico")
    st.info(f"Período selecionado: **{month_name}/{year}**")
    
    if 'mostrar_importacao' not in st.session_state: st.session_state.mostrar_importacao = False
    if 'mostrar_confirmacao_limpar' not in st.session_state: st.session_state.mostrar_confirmacao_limpar = False
    if 'importacao_em_andamento' not in st.session_state: st.session_state.importacao_em_andamento = False
    
    df_horasponto = get_horasponto_by_periodo(year, month_num)
    pode_editar = tipo_usuario in ['Admin', 'Gestor']
    
    if pode_editar:
        col_btn1, col_btn2 = st.columns([1, 1])
        with col_btn1:
            if st.button("📤 Importar CSV", width="stretch", type="primary"):
                st.session_state.mostrar_importacao = True
                st.session_state.mostrar_confirmacao_limpar = False
                st.session_state.importacao_em_andamento = False
                st.rerun()
        with col_btn2:
            if not df_horasponto.empty:
                if st.button("🗑️ Limpar Período", width="stretch", type="secondary"):
                    st.session_state.mostrar_confirmacao_limpar = True
                    st.session_state.mostrar_importacao = False
                    st.rerun()
        
        if st.session_state.get('mostrar_confirmacao_limpar'):
            st.warning(f"⚠️ Você está prestes a excluir **todos os registros** de {month_name}/{year}")
            st.info(f"Total de registros: **{len(df_horasponto)}**")
            col_sim, col_nao = st.columns(2)
            with col_sim:
                if st.button("✅ Sim, Excluir", width="stretch", type="primary"):
                    sucesso, msg = deletar_horasponto_by_periodo(year, month_num)
                    if sucesso:
                        st.success(f"✅ {msg}")
                        st.session_state.mostrar_confirmacao_limpar = False
                        st.rerun()
                    else: st.error(f"❌ {msg}")
            with col_nao:
                if st.button("❌ Não, Cancelar", width="stretch"):
                    st.session_state.mostrar_confirmacao_limpar = False
                    st.rerun()
        
        if st.session_state.get('mostrar_importacao'):
            st.markdown("---")
            st.subheader("📤 Importar Arquivo CSV de Ponto")
            
            if not df_horasponto.empty:
                st.warning(f"⚠️ **Atenção:** Já existem **{len(df_horasponto)} registros** para {month_name}/{year}")
                st.info("Os dados existentes serão **substituídos** pelos dados do arquivo CSV.")
            
            uploaded_file = st.file_uploader(
                "Selecione o arquivo CSV do ponto",
                type=['csv'],
                help="O arquivo deve conter as colunas: colaborador_ponto, horas_ponto, afastamento"
            )
            
            if uploaded_file is not None:
                try:
                    raw_bytes = uploaded_file.read()
                    uploaded_file.seek(0)
                    
                    separador = ','
                    try:
                        texto_amostra = raw_bytes[:2048].decode('utf-8', errors='ignore')
                        primeira_linha = texto_amostra.split('\n')[0]
                        if primeira_linha.count(';') > primeira_linha.count(','): separador = ';'
                    except: pass
                    
                    try: df_csv = pd.read_csv(io.BytesIO(raw_bytes), encoding='utf-8', sep=separador)
                    except UnicodeDecodeError: df_csv = pd.read_csv(io.BytesIO(raw_bytes), encoding='latin1', sep=separador)
                    
                    df_csv.columns = [col.strip().lower() for col in df_csv.columns]
                    colunas_necessarias = ['colaborador_ponto', 'horas_ponto', 'afastamento']
                    colunas_faltantes = [col for col in colunas_necessarias if col not in df_csv.columns]
                    
                    if colunas_faltantes:
                        st.error(f"❌ Colunas faltantes no CSV: {', '.join(colunas_faltantes)}")
                    else:
                        st.success(f"✅ Arquivo carregado: **{len(df_csv)} registros** encontrados")
                        st.dataframe(df_csv.head(10), width="stretch")
                        
                        col_confirmar, col_cancelar = st.columns(2)
                        with col_confirmar:
                            if st.button("✅ Confirmar Importação", width="stretch", type="primary", key="btn_confirmar_importacao"):
                                st.session_state.importacao_em_andamento = True
                                st.rerun()
                        
                        with col_cancelar:
                            if st.button("❌ Cancelar", width="stretch"):
                                st.session_state.mostrar_importacao = False
                                st.session_state.importacao_em_andamento = False
                                st.rerun()
                        
                        if st.session_state.get('importacao_em_andamento'):
                            with st.spinner("🔄 Processando importação em lote otimizado..."):
                                colaboradores_map = get_colaboradores_map()
                                
                                if not df_horasponto.empty:
                                    deletar_horasponto_by_periodo(year, month_num)
                                
                                registros_para_inserir = []
                                colaboradores_nao_encontrados = []
                                
                                for _, row in df_csv.iterrows():
                                    nome_ponto = str(row['colaborador_ponto']).strip()
                                    try: horas = float(str(row['horas_ponto']).replace(',', '.')) if pd.notna(row['horas_ponto']) else 0.0
                                    except: horas = 0.0
                                    try: afastamento = float(str(row['afastamento']).replace(',', '.')) if pd.notna(row['afastamento']) else 0.0
                                    except: afastamento = 0.0
                                    
                                    id_colab = encontrar_colaborador_id(nome_ponto, colaboradores_map)
                                    if id_colab:
                                        registros_para_inserir.append({
                                            'ano': year,
                                            'mes': month_num,
                                            'id_colaborador': id_colab,
                                            'colaborador_ponto': nome_ponto,
                                            'horas_ponto': horas,
                                            'afastamento': afastamento
                                        })
                                    else:
                                        colaboradores_nao_encontrados.append(nome_ponto)
                                
                                inseridos, erros, err_list = inserir_horasponto_lote(registros_para_inserir)
                                st.session_state.importacao_em_andamento = False
                                
                                st.markdown("---")
                                st.markdown("### 📊 Resultado da Importação")
                                col1, col2, col3 = st.columns(3)
                                col1.metric("✅ Registros Inseridos", inseridos)
                                col2.metric("❌ Não Encontrados", len(colaboradores_nao_encontrados))
                                col3.metric("📋 Total Processados", len(df_csv))
                                
                                if colaboradores_nao_encontrados:
                                    st.warning(f"⚠️ **{len(colaboradores_nao_encontrados)} colaborador(es) não encontrados** no cadastro:")
                                    st.dataframe(pd.DataFrame({'Nome no CSV': colaboradores_nao_encontrados}), width="stretch")
                                
                                if inseridos > 0:
                                    st.success(f"🎉 **Importação concluída!** {inseridos} registros gravados com sucesso.")
                                
                                st.session_state.mostrar_importacao = False
                                st.rerun()
                except Exception as e:
                    st.error(f"❌ Erro ao processar arquivo: {str(e)}")
                    st.session_state.importacao_em_andamento = False
    
    st.markdown("---")
    st.markdown("### 📋 Registros do Período")
    df_horasponto_atualizado = get_horasponto_by_periodo(year, month_num)
    
    if df_horasponto_atualizado.empty:
        st.info(f"📭 Nenhum registro de ponto encontrado para **{month_name}/{year}**")
    else:
        st.markdown(f"**Total de registros:** {len(df_horasponto_atualizado)}")
        colunas_exibicao = [c for c in ['colaborador_ponto', 'cargo', 'nivel', 'horas_ponto', 'afastamento'] if c in df_horasponto_atualizado.columns]
        st.dataframe(
            df_horasponto_atualizado[colunas_exibicao],
            width="stretch",
            hide_index=True,
            column_config={
                "colaborador_ponto": st.column_config.TextColumn("Colaborador", width="large"),
                "cargo": st.column_config.TextColumn("Cargo"),
                "nivel": st.column_config.TextColumn("Nível"),
                "horas_ponto": st.column_config.NumberColumn("Horas Ponto", format="%.2f h"),
                "afastamento": st.column_config.NumberColumn("Afastamento", format="%.2f h")
            }
        )
        col_total1, col_total2 = st.columns(2)
        col_total1.metric("🕐 Total de Horas", f"{df_horasponto_atualizado['horas_ponto'].sum():.2f} h")
        col_total2.metric("🏖️ Total de Afastamento", f"{df_horasponto_atualizado['afastamento'].sum():.2f} h")
        
        st.markdown("---")
        csv_data = df_horasponto_atualizado.to_csv(index=False, encoding='utf-8')
        st.download_button("⬇️ Baixar CSV do Período", data=csv_data, file_name=f"horasponto_{year}_{month_name}.csv", mime="text/csv", width="stretch")

# =============================================================================
# CRUD DE PROJETOS
# =============================================================================
def render_projetos_tab():
    st.subheader("🏢 Gerenciamento de Projetos")
    if 'projeto_editando' not in st.session_state: st.session_state.projeto_editando = None
    if 'mostrar_form_projeto' not in st.session_state: st.session_state.mostrar_form_projeto = False
    if 'mostrar_confirmacao_exclusao_projeto' not in st.session_state: st.session_state.mostrar_confirmacao_exclusao_projeto = None
    
    col_btn1, col_btn2 = st.columns([3, 1])
    with col_btn2:
        if st.button("➕ Novo Projeto", width="stretch", type="primary"):
            st.session_state.projeto_editando = None; st.session_state.mostrar_form_projeto = True; st.rerun()
    
    df_projetos = listar_projetos()
    if df_projetos.empty:
        st.warning("⚠️ Nenhum projeto cadastrado. Clique em **➕ Novo Projeto** para começar.")
    else:
        st.markdown(f"**Total de projetos cadastrados:** {len(df_projetos)}")
        df_exibicao = df_projetos.copy()
        for col in ['created_at', 'criado_em', 'updated_at', 'atualizado_em']:
            if col in df_exibicao.columns: df_exibicao[col] = df_exibicao[col].apply(lambda x: x.strftime("%d/%m/%Y %H:%M") if pd.notna(x) else "")
        if 'is_active' in df_exibicao.columns:
            df_exibicao['status'] = df_exibicao['is_active'].apply(lambda x: "✅ Ativo" if x else "❌ Inativo")
        
        colunas_exibicao = [c for c in ['id', 'name', 'status', 'created_at', 'criado_em', 'updated_at', 'atualizado_em'] if c in df_exibicao.columns]
        st.dataframe(df_exibicao[colunas_exibicao], width="stretch", hide_index=True,
            column_config={"id": st.column_config.NumberColumn("ID", width="small"), "name": st.column_config.TextColumn("Nome do Projeto", width="large"), "status": st.column_config.TextColumn("Status")})
        
        st.markdown("---")
        st.markdown("### Ações")
        options = [f"{row['id']} - {row['name']} ({'Ativo' if row.get('is_active') else 'Inativo'})" for _, row in df_projetos.iterrows()]
        projeto_selecionado = st.selectbox("Selecione um projeto para editar ou excluir:", options=options, index=0)
        if projeto_selecionado:
            projeto_id = int(projeto_selecionado.split(" - ")[0])
            projeto_dados = df_projetos[df_projetos['id'] == projeto_id].iloc[0].to_dict()
            col_edit, col_del = st.columns(2)
            with col_edit:
                if st.button("✏️ Editar Projeto", width="stretch"): st.session_state.projeto_editando = projeto_dados; st.session_state.mostrar_form_projeto = True; st.rerun()
            with col_del:
                if st.button("🗑️ Excluir Projeto", width="stretch", type="secondary"): st.session_state.mostrar_confirmacao_exclusao_projeto = projeto_id; st.rerun()
    
    if st.session_state.get('mostrar_confirmacao_exclusao_projeto'):
        col_id = st.session_state.mostrar_confirmacao_exclusao_projeto
        col_nome = df_projetos[df_projetos['id'] == col_id]['name'].values[0]
        st.warning(f"⚠️ Você está prestes a excluir o projeto: **{col_nome}**")
        st.info("Esta ação não pode ser desfeita!")
        col_sim, col_nao = st.columns(2)
        with col_sim:
            if st.button("✅ Sim, Excluir", width="stretch", type="primary"):
                sucesso, msg = deletar_projeto(col_id)
                if sucesso: st.success(f"✅ {msg}"); st.session_state.mostrar_confirmacao_exclusao_projeto = None; st.rerun()
                else: st.error(f"❌ {msg}")
        with col_nao:
            if st.button("❌ Não, Cancelar", width="stretch"): st.session_state.mostrar_confirmacao_exclusao_projeto = None; st.rerun()
    
    if st.session_state.get('mostrar_form_projeto'):
        st.markdown("---")
        st.subheader("📝 " + ("Editar Projeto" if st.session_state.projeto_editando else "Novo Projeto"))
        projeto = st.session_state.projeto_editando or {}
        with st.form("form_projeto"):
            name = st.text_input("Nome do Projeto *", value=projeto.get('name', ''), placeholder="Ex: ASCOM - Comunicação")
            is_active = st.checkbox("Projeto Ativo", value=projeto.get('is_active', True))
            st.markdown("---")
            col_submit, col_cancel = st.columns(2)
            with col_submit: submitted = st.form_submit_button("💾 Salvar", width="stretch", type="primary")
            with col_cancel: cancelled = st.form_submit_button("❌ Cancelar", width="stretch")
            if cancelled: st.session_state.projeto_editando = None; st.session_state.mostrar_form_projeto = False; st.rerun()
            if submitted:
                if not name.strip(): st.error("⚠️ Nome do projeto é obrigatório")
                else:
                    if st.session_state.projeto_editando: sucesso, msg = atualizar_projeto(st.session_state.projeto_editando['id'], name, is_active)
                    else: sucesso, msg = criar_projeto(name, is_active)
                    if sucesso: st.success(f"✅ {msg}"); st.session_state.projeto_editando = None; st.session_state.mostrar_form_projeto = False; st.rerun()
                    else: st.error(f"❌ {msg}")

# =============================================================================
# MÉTRICAS
# =============================================================================
# =============================================================================
# MÉTRICAS E INDICADORES ESTRATÉGICOS (COMPLETO: SECÇÕES 1 A 5)
# =============================================================================
def render_metricas_tab(all_workitems, project_of_item, df_colaboradores, df_horasponto, df_valores, dias_uteis, year, month_name, month_num, selected_wi_types, selected_states, parents_data):
    st.subheader("📊 Métricas e Indicadores Estratégicos")
    dias_uteis_reais = dias_uteis / 8.0 if dias_uteis > 0 else 0
    st.info(f"Período: **{month_name}/{year}** | Dias úteis: **{dias_uteis_reais:.0f}** ({dias_uteis:.0f}h)")
    
    # 1. Leitura em cache de capacities e valores por nível
    all_capacities = get_all_capacities()
    nivel_valor = {
        str(r["nivel"]).strip(): limpar_valor_monetario(r.get("valor")) 
        for _, r in df_valores.iterrows() if str(r.get("nivel", "")).strip()
    }
    
    df_colab_periodo = df_colaboradores.copy()
    azure_to_ponto, colab_nivel_map, colab_cargo_map = {}, {}, {}
    for _, row in df_colab_periodo.iterrows():
        azure_name = str(row["colaborador_azure"]).strip().lower()
        ponto_name = str(row["colaborador_ponto"]).strip()
        if azure_name: 
            azure_to_ponto[azure_name] = ponto_name
        if ponto_name:
            colab_nivel_map[ponto_name] = str(row.get("nivel", "")).strip()
            colab_cargo_map[ponto_name] = str(row.get("cargo", "")).strip()
    
    # 2. Consolidação de horas do Azure para colaboradores cadastrados
    horas_azure_por_colab = {}
    for item_id, item in all_workitems.items():
        f = item.get("fields", {})
        if f.get("System.State") != "Done": 
            continue
        assigned_name = get_assigned_to_name(f.get("System.AssignedTo"))
        colaborador_ponto = encontrar_colaborador_ponto(assigned_name, azure_to_ponto)
        if colaborador_ponto != "Não atribuído" and colaborador_ponto in colab_nivel_map:
            horas_azure_por_colab[colaborador_ponto] = horas_azure_por_colab.get(colaborador_ponto, 0.0) + float(f.get("Microsoft.VSTS.Scheduling.CompletedWork") or 0.0)
    
    # -------------------------------------------------------------------------
    # SECÇÃO 1: MÉTRICAS FINANCEIRAS E DE EFICIÊNCIA
    # -------------------------------------------------------------------------
    st.markdown("---")
    st.markdown("### 💰 1. Métricas Financeiras e de Eficiência")
    total_capacity_geral, total_horas_ponto, total_horas_azure, total_horas_afastamento, total_faturado = 0.0, 0.0, 0.0, 0.0, 0.0
    metricas_colab = []
    
    for _, hp_row in df_horasponto.iterrows():
        ponto_name = str(hp_row["colaborador_ponto"]).strip()
        if ponto_name not in colab_nivel_map: 
            continue
        
        horas_ponto = float(hp_row["horas_ponto"]) if pd.notna(hp_row.get("horas_ponto")) else 0.0
        afastamento = float(hp_row["afastamento"]) if pd.notna(hp_row.get("afastamento")) else 0.0
        horas_azure = horas_azure_por_colab.get(ponto_name, 0.0)
        
        nivel = colab_nivel_map[ponto_name]
        cargo = colab_cargo_map.get(ponto_name, "")
        capacity = all_capacities.get(nivel, 0.0)
        
        capacity_calculado = dias_uteis_reais * capacity if capacity > 0 else 0.0
        utilizacao = (horas_azure / capacity_calculado * 100) if capacity_calculado > 0 else 0.0
        gap_apontamento = horas_ponto - horas_azure
        capacidade_ociosa = capacity_calculado - horas_azure
        
        valor_profissional = nivel_valor.get(nivel, 0.0)
        valor_hora = valor_profissional / dias_uteis if dias_uteis > 0 else 0.0
        valor_recalculado = round(horas_ponto * valor_hora, 2)
        
        metricas_colab.append({
            "Colaborador": ponto_name, 
            "Cargo": cargo, 
            "Nível": nivel,
            "Horas Ponto": horas_ponto, 
            "Horas Azure": horas_azure,
            "Capacity": capacity_calculado, 
            "Utilização %": round(utilizacao, 1),
            "Gap Apontamento": round(gap_apontamento, 2), 
            "Capacidade Ociosa": round(capacidade_ociosa, 2),
            "Afastamento": afastamento, 
            "Faturamento": valor_recalculado
        })
        
        total_capacity_geral += capacity_calculado
        total_horas_ponto += horas_ponto
        total_horas_azure += horas_azure
        total_horas_afastamento += afastamento
        total_faturado += valor_recalculado
    
    utilizacao_geral = (total_horas_azure / total_capacity_geral * 100) if total_capacity_geral > 0 else 0.0
    gap_geral = total_horas_ponto - total_horas_azure
    capacidade_ociosa_geral = total_capacity_geral - total_horas_azure
    ticket_medio_hora = total_faturado / total_horas_ponto if total_horas_ponto > 0 else 0.0
    
    col1, col2, col3, col4 = st.columns(4)
    col1.metric("📊 Utilização da Capacidade", f"{utilizacao_geral:.1f}%")
    col2.metric("⚠️ Gap de Apontamento", f"{gap_geral:.1f} h", help="Horas trabalhadas (Ponto) - Horas registradas no Azure")
    col3.metric("💵 Ticket Médio/Hora", format_moeda_brl(ticket_medio_hora))
    col4.metric("💰 Total Faturado", format_moeda_brl(total_faturado))
    
    col5, col6 = st.columns(2)
    col5.metric("📉 Capacidade Ociosa", f"{capacidade_ociosa_geral:.1f} h", help="Capacity teórico - Horas Azure")
    col6.metric("🏖️ Total Afastamento", f"{total_horas_afastamento:.1f} h")
    
    if metricas_colab:
        df_metricas = pd.DataFrame(metricas_colab)
        st.markdown("---")
        st.subheader("📈 Utilização da Capacidade por Colaborador")
        fig_util = px.bar(
            df_metricas, 
            x="Colaborador", 
            y="Utilização %", 
            color="Utilização %", 
            color_continuous_scale=["#dc3545", "#ffc107", "#28a745"], 
            text="Utilização %"
        )
        fig_util.update_layout(margin=dict(t=20, b=80, l=20, r=20), height=400, xaxis_tickangle=-45)
        fig_util.update_traces(texttemplate='%{text:.0f}%', textposition='outside')
        fig_util.add_hline(y=85, line_dash="dash", line_color="green", annotation_text="Ideal (85%)")
        fig_util.add_hline(y=100, line_dash="dash", line_color="red", annotation_text="Máximo (100%)")
        st.plotly_chart(fig_util, width="stretch")
        
        if total_horas_ponto > 0 and gap_geral > total_horas_ponto * 0.1:
            st.warning(f"⚠️ **Alerta:** Gap de apontamento significativo ({gap_geral:.1f}h = {gap_geral/total_horas_ponto*100:.1f}% das horas do ponto não registradas no Azure)")
        if total_capacity_geral > 0 and capacidade_ociosa_geral > total_capacity_geral * 0.2:
            st.warning(f"⚠️ **Alerta:** Capacidade ociosa alta ({capacidade_ociosa_geral:.1f}h = {capacidade_ociosa_geral/total_capacity_geral*100:.1f}% da capacidade não utilizada)")

    # -------------------------------------------------------------------------
    # SECÇÃO 2: MÉTRICAS DE PRODUTIVIDADE E AGILE
    # -------------------------------------------------------------------------
    st.markdown("---")
    st.markdown("### 🚀 2. Métricas de Produtividade e Agile")
    lead_times, cycle_times, items_done = [], [], []
    for item_id, item in all_workitems.items():
        f = item.get("fields", {})
        if f.get("System.State") != "Done": 
            continue
        created = parse_date(f.get("System.CreatedDate"))
        closed = parse_date(f.get("Microsoft.VSTS.Common.ClosedDate"))
        activated = parse_date(f.get("Microsoft.VSTS.Common.ActivatedDate"))
        if created and closed: 
            lead_times.append((closed - created).total_seconds() / 3600)
        if activated and closed: 
            cycle_times.append((closed - activated).total_seconds() / 3600)
        items_done.append(item)
    
    lead_time_medio = sum(lead_times) / len(lead_times) if lead_times else 0
    cycle_time_medio = sum(cycle_times) / len(cycle_times) if cycle_times else 0
    
    sprints_dict = {}
    for item in items_done:
        f = item.get("fields", {})
        sprint_name = str(f.get("System.IterationPath", "N/A")).split("\\")[-1]
        if sprint_name not in sprints_dict: 
            sprints_dict[sprint_name] = {"qtd": 0, "horas": 0.0}
        sprints_dict[sprint_name]["qtd"] += 1
        sprints_dict[sprint_name]["horas"] += float(f.get("Microsoft.VSTS.Scheduling.CompletedWork") or 0.0)
    
    velocidade_media_itens = sum(s["qtd"] for s in sprints_dict.values()) / len(sprints_dict) if sprints_dict else 0
    velocidade_media_horas = sum(s["horas"] for s in sprints_dict.values()) / len(sprints_dict) if sprints_dict else 0
    
    col1, col2, col3, col4 = st.columns(4)
    col1.metric("⏱️ Lead Time Médio", f"{lead_time_medio/24:.1f} dias")
    col2.metric("🔄 Cycle Time Médio", f"{cycle_time_medio/24:.1f} dias")
    col3.metric("📦 Velocidade Média", f"{velocidade_media_itens:.0f} itens/sprint")
    col4.metric("⚡ Velocidade em Horas", f"{velocidade_media_horas:.1f}h/sprint")
    
    if sprints_dict:
        st.markdown("---")
        st.subheader("📦 Velocidade por Sprint")
        df_velocidade = pd.DataFrame([{"Sprint": k, "Itens": v["qtd"], "Horas": round(v["horas"], 2)} for k, v in sprints_dict.items()])
        df_velocidade = df_velocidade.sort_values("Horas", ascending=False).reset_index(drop=True)
        
        fig_vel = go.Figure()
        fig_vel.add_trace(go.Bar(
            x=df_velocidade["Sprint"], 
            y=df_velocidade["Itens"], 
            name="Itens", 
            marker_color="#2C5282", 
            text=df_velocidade["Itens"], 
            textposition='outside', 
            textfont=dict(size=10)
        ))
        fig_vel.add_trace(go.Scatter(
            x=df_velocidade["Sprint"], 
            y=df_velocidade["Horas"], 
            name="Horas", 
            mode='lines+markers', 
            marker=dict(color="#FF6600", size=8), 
            line=dict(color="#FF6600", width=3), 
            yaxis="y2", 
            text=df_velocidade["Horas"].apply(lambda x: f"{x:.0f}h"), 
            textposition='top center', 
            textfont=dict(size=10, color="#FF6600")
        ))
        fig_vel.update_layout(
            title="Itens e Horas por Sprint (ordenado por horas)", 
            xaxis=dict(title="Sprint", tickangle=-45), 
            yaxis=dict(title="Quantidade de Itens", side="left"), 
            yaxis2=dict(title="Horas Apontadas", side="right", overlaying="y", showgrid=False), 
            margin=dict(t=60, b=80, l=60, r=60), 
            height=450, 
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1), 
            barmode='group'
        )
        st.plotly_chart(fig_vel, width="stretch")

    # -------------------------------------------------------------------------
    # SECÇÃO 3: DISTRIBUIÇÃO DE ESFORÇO: MANUTENÇÃO VS EVOLUÇÃO
    # -------------------------------------------------------------------------
    st.markdown("---")
    st.markdown("### 🎯 3. Distribuição de Esforço: Manutenção vs Evolução")
    categorias = {
        "Evolução (Novas Funcionalidades)": ["Product Backlog Item", "Melhoria"], 
        "Manutenção Corretiva": ["Bug", "BugFix", "HotFix"], 
        "Débito Técnico": ["Débito Técnico"], 
        "Outros": ["Task", "Design"]
    }
    horas_por_categoria = {cat: 0.0 for cat in categorias.keys()}
    qtd_por_categoria = {cat: 0 for cat in categorias.keys()}
    
    for item_id, item in all_workitems.items():
        f = item.get("fields", {})
        wtype = f.get("System.WorkItemType", "")
        if wtype not in selected_wi_types: 
            continue
        completed = float(f.get("Microsoft.VSTS.Scheduling.CompletedWork") or 0.0)
        for cat, tipos in categorias.items():
            if wtype in tipos: 
                horas_por_categoria[cat] += completed
                qtd_por_categoria[cat] += 1
                break
    
    total_horas_esforco = sum(horas_por_categoria.values())
    col1, col2 = st.columns(2)
    with col1:
        st.markdown("#### 📊 Distribuição por Horas Apontadas")
        if total_horas_esforco > 0:
            df_esforco = pd.DataFrame([
                {"Categoria": cat, "Horas": round(horas_por_categoria[cat], 2), "Percentual": round(horas_por_categoria[cat]/total_horas_esforco*100, 1)} 
                for cat in categorias.keys()
            ])
            fig_esforco = px.pie(df_esforco, values="Horas", names="Categoria", hole=0.5, color_discrete_sequence=px.colors.qualitative.Set2)
            fig_esforco.update_traces(textposition='inside', textinfo='percent+label')
            fig_esforco.update_layout(margin=dict(t=20, b=20, l=20, r=20), height=350)
            st.plotly_chart(fig_esforco, width="stretch")
        else: 
            st.info("Sem dados de horas apontadas")
            
    with col2:
        st.markdown("#### 📋 Distribuição por Quantidade de Itens")
        total_qtd = sum(qtd_por_categoria.values())
        if total_qtd > 0:
            df_qtd = pd.DataFrame([
                {"Categoria": cat, "Quantidade": qtd_por_categoria[cat], "Percentual": round(qtd_por_categoria[cat]/total_qtd*100, 1)} 
                for cat in categorias.keys()
            ])
            fig_qtd = px.pie(df_qtd, values="Quantidade", names="Categoria", hole=0.5, color_discrete_sequence=px.colors.qualitative.Pastel)
            fig_qtd.update_traces(textposition='inside', textinfo='percent+label')
            fig_qtd.update_layout(margin=dict(t=20, b=20, l=20, r=20), height=350)
            st.plotly_chart(fig_qtd, width="stretch")
        else: 
            st.info("Sem dados de quantidade")
    
    manutencao_pct = (horas_por_categoria["Manutenção Corretiva"] / total_horas_esforco * 100) if total_horas_esforco > 0 else 0
    debito_pct = (horas_por_categoria["Débito Técnico"] / total_horas_esforco * 100) if total_horas_esforco > 0 else 0
    if manutencao_pct + debito_pct > 40: 
        st.error(f"🚨 **Alerta Crítico:** {(manutencao_pct + debito_pct):.0f}% do esforço está em manutenção/débito técnico (ideal: < 30%)")
    elif manutencao_pct + debito_pct > 30: 
        st.warning(f"⚠️ **Atenção:** {(manutencao_pct + debito_pct):.0f}% do esforço está em manutenção/débito técnico")
    else: 
        st.success(f"✅ **Saudável:** Apenas {(manutencao_pct + debito_pct):.0f}% em manutenção/débito técnico")

    # -------------------------------------------------------------------------
    # SECÇÃO 4: MÉTRICAS DE GESTÃO DE PESSOAS
    # -------------------------------------------------------------------------
    st.markdown("---")
    st.markdown("### 👥 4. Métricas de Gestão de Pessoas")
    taxa_absenteismo = (total_horas_afastamento / total_capacity_geral * 100) if total_capacity_geral > 0 else 0.0
    
    top_contributors = []
    for item_id, item in all_workitems.items():
        f = item.get("fields", {})
        if f.get("System.State") != "Done": 
            continue
        top_contributors.append({
            "Colaborador": get_assigned_to_name(f.get("System.AssignedTo")), 
            "Horas": float(f.get("Microsoft.VSTS.Scheduling.CompletedWork") or 0.0)
        })
    
    df_top = pd.DataFrame(top_contributors)
    
    st.markdown("#### 🏖️ Taxa de Absenteísmo")
    st.metric(
        "Taxa de Afastamento", 
        f"{taxa_absenteismo:.1f}%", 
        help=f"Horas de afastamento ({total_horas_afastamento:.0f}h) / Capacity total ({total_capacity_geral:.0f}h)"
    )
    if taxa_absenteismo > 15: 
        st.error(f"🚨 Absenteísmo alto ({taxa_absenteismo:.1f}%). Pode comprometer entregas.")
    elif taxa_absenteismo > 10: 
        st.warning(f"⚠️ Absenteísmo moderado ({taxa_absenteismo:.1f}%)")
    else: 
        st.success(f"✅ Absenteísmo dentro do esperado ({taxa_absenteismo:.1f}%)")
    
    if not df_top.empty:
        df_top_agg = df_top.groupby("Colaborador").agg({"Horas": "sum"}).reset_index()
        df_top_10 = df_top_agg.sort_values("Horas", ascending=False).head(10)
        df_bottom_10 = df_top_agg.sort_values("Horas", ascending=True).head(10)
        
        st.markdown("---")
        col_top, col_bottom = st.columns(2)
        with col_top:
            st.markdown("#### 🏆 Top 10 Contributors (Mais Horas)")
            fig_top = px.bar(df_top_10, x="Horas", y="Colaborador", orientation='h', color="Horas", color_continuous_scale="Greens")
            fig_top.update_traces(texttemplate='%{x:.0f}h', textposition='outside')
            fig_top.update_layout(margin=dict(t=20, b=20, l=150, r=20), height=400, showlegend=False)
            st.plotly_chart(fig_top, width="stretch")
        
        with col_bottom:
            st.markdown("#### ⚠️ Bottom 10 Contributors (Menos Horas)")
            fig_bottom = px.bar(df_bottom_10, x="Horas", y="Colaborador", orientation='h', color="Horas", color_continuous_scale="Reds")
            fig_bottom.update_traces(texttemplate='%{x:.0f}h', textposition='outside')
            fig_bottom.update_layout(margin=dict(t=20, b=20, l=150, r=20), height=400, showlegend=False)
            st.plotly_chart(fig_bottom, width="stretch")
    else:
        st.info("Sem dados de contribuição")

    # -------------------------------------------------------------------------
    # SECÇÃO 5: RESUMO EXECUTIVO DO PERÍODO
    # -------------------------------------------------------------------------
    st.markdown("---")
    st.markdown("### 📋 5. Resumo Executivo do Período")
    pct_evolucao = (horas_por_categoria['Evolução (Novas Funcionalidades)'] / total_horas_esforco * 100) if total_horas_esforco > 0 else 0.0
    pct_manutencao = (horas_por_categoria['Manutenção Corretiva'] / total_horas_esforco * 100) if total_horas_esforco > 0 else 0.0
    
    resumo_data = {
        "Indicador": [
            "Total de Colaboradores no Período", "Total de Horas Apontadas (Azure)", "Total de Horas do Ponto", 
            "Capacidade Total da Equipe", "Taxa de Utilização", "Gap de Apontamento (Ponto vs Azure)", 
            "Capacidade Ociosa (Capacity vs Azure)", "Total Faturado", "Ticket Médio por Hora", 
            "Lead Time Médio", "Cycle Time Médio", "Velocidade (itens/sprint)", 
            "% Manutenção/Corretivo", "% Evolução/Novas Func.", "Taxa de Absenteísmo"
        ],
        "Valor": [
            f"{len(df_colab_periodo)}", f"{total_horas_azure:.1f} h", f"{total_horas_ponto:.1f} h", 
            f"{total_capacity_geral:.1f} h", f"{utilizacao_geral:.1f}%", f"{gap_geral:.1f} h", 
            f"{capacidade_ociosa_geral:.1f} h", format_moeda_brl(total_faturado), format_moeda_brl(ticket_medio_hora), 
            f"{lead_time_medio/24:.1f} dias", f"{cycle_time_medio/24:.1f} dias", f"{velocidade_media_itens:.0f} itens", 
            f"{pct_manutencao:.1f}%", f"{pct_evolucao:.1f}%", f"{taxa_absenteismo:.1f}%"
        ],
        "Fórmula": [
            "Count(colaboradores)", "SUM(CompletedWork)", "SUM(horas_ponto)", 
            "Dias_úteis × Capacity_nível", "Azure / Capacity × 100", "Ponto - Azure", 
            "Capacity - Azure", "Horas × Valor_hora", "Faturado / Horas_ponto", 
            "ClosedDate - CreatedDate", "ClosedDate - ActivatedDate", "Itens Done / Sprints", 
            "Manutenção / Total × 100", "Evolução / Total × 100", "Afastamento / Capacity × 100"
        ],
        "Status": [
            "✅" if len(df_colab_periodo) > 0 else "⚠️", 
            "✅" if total_horas_azure > 0 else "⚠️", 
            "✅" if total_horas_ponto > 0 else "⚠️", 
            "✅" if total_capacity_geral > 0 else "⚠️", 
            "✅" if 85 <= utilizacao_geral <= 95 else ("⚠️" if utilizacao_geral > 0 else "❌"), 
            "✅" if abs(gap_geral) < total_horas_ponto * 0.05 else "⚠️", 
            "✅" if capacidade_ociosa_geral < total_capacity_geral * 0.2 else "⚠️", 
            "✅" if total_faturado > 0 else "⚠️", 
            "✅" if ticket_medio_hora > 0 else "⚠️", 
            "✅" if lead_time_medio > 0 else "⚠️", 
            "✅" if cycle_time_medio > 0 else "⚠️", 
            "✅" if velocidade_media_itens > 0 else "⚠️", 
            "✅" if pct_manutencao < 30 else ("⚠️" if pct_manutencao < 40 else "❌"), 
            "✅" if pct_evolucao > 50 else "⚠️", 
            "✅" if taxa_absenteismo < 10 else ("⚠️" if taxa_absenteismo < 15 else "❌")
        ]
    }
    st.dataframe(
        pd.DataFrame(resumo_data), 
        width="stretch", 
        hide_index=True, 
        column_config={
            "Indicador": st.column_config.TextColumn("Indicador", width="large"), 
            "Valor": st.column_config.TextColumn("Valor"), 
            "Fórmula": st.column_config.TextColumn("Fórmula de Cálculo", width="large"), 
            "Status": st.column_config.TextColumn("Status")
        }
    )
# =============================================================================
# CRUD COLABORADORES - RENDER
# =============================================================================
def render_colaboradores_tab():
    st.subheader("⚙️ Gerenciamento de Colaboradores")
    if 'colaborador_editando' not in st.session_state: st.session_state.colaborador_editando = None
    if 'mostrar_form_colaborador' not in st.session_state: st.session_state.mostrar_form_colaborador = False
    if 'mostrar_confirmacao_exclusao' not in st.session_state: st.session_state.mostrar_confirmacao_exclusao = None
    
    col_btn1, col_btn2 = st.columns([3, 1])
    with col_btn2:
        if st.button("➕ Novo Colaborador", width="stretch", type="primary"):
            st.session_state.colaborador_editando = None; st.session_state.mostrar_form_colaborador = True; st.rerun()
    
    df_colaboradores = listar_colaboradores()
    if df_colaboradores.empty:
        st.warning("⚠️ Nenhum colaborador cadastrado. Clique em **➕ Novo Colaborador** para começar.")
    else:
        st.markdown(f"**Total de colaboradores cadastrados:** {len(df_colaboradores)}")
        df_exibicao = df_colaboradores.copy()
        for col in ['data_inicio', 'data_saida', 'promocao']:
            if col in df_exibicao.columns: df_exibicao[col] = df_exibicao[col].apply(lambda x: x.strftime("%d/%m/%Y") if pd.notna(x) else "")
        
        colunas_exibicao = [c for c in ['id', 'colaborador_ponto', 'cargo', 'nivel', 'email', 'data_inicio', 'promocao', 'reposicao', 'data_saida', 'status'] if c in df_exibicao.columns]
        st.dataframe(df_exibicao[colunas_exibicao], width="stretch", hide_index=True,
            column_config={"id": st.column_config.NumberColumn("ID", width="small"), "colaborador_ponto": st.column_config.TextColumn("Nome", width="large")})
        
        st.markdown("---")
        st.markdown("### Ações")
        colaborador_selecionado = st.selectbox("Selecione um colaborador para editar ou excluir:", options=[f"{row['id']} - {row['colaborador_ponto']} ({row['cargo']})" for _, row in df_colaboradores.iterrows()], index=0)
        if colaborador_selecionado:
            colaborador_id = int(colaborador_selecionado.split(" - ")[0])
            colaborador_dados = df_colaboradores[df_colaboradores['id'] == colaborador_id].iloc[0].to_dict()
            col_edit, col_del = st.columns(2)
            with col_edit:
                if st.button("✏️ Editar Colaborador", width="stretch"): st.session_state.colaborador_editando = colaborador_dados; st.session_state.mostrar_form_colaborador = True; st.rerun()
            with col_del:
                if st.button("🗑️ Excluir Colaborador", width="stretch", type="secondary"): st.session_state.mostrar_confirmacao_exclusao = colaborador_id; st.rerun()
    
    if st.session_state.get('mostrar_confirmacao_exclusao'):
        col_id = st.session_state.mostrar_confirmacao_exclusao
        col_nome = df_colaboradores[df_colaboradores['id'] == col_id]['colaborador_ponto'].values[0]
        st.warning(f"⚠️ Você está prestes a excluir o colaborador: **{col_nome}**")
        st.info("Esta ação não pode ser desfeita!")
        col_sim, col_nao = st.columns(2)
        with col_sim:
            if st.button("✅ Sim, Excluir", width="stretch", type="primary"):
                sucesso, msg = deletar_colaborador(col_id)
                if sucesso: st.success(f"✅ {msg}"); st.session_state.mostrar_confirmacao_exclusao = None; st.rerun()
                else: st.error(f"❌ {msg}")
        with col_nao:
            if st.button("❌ Não, Cancelar", width="stretch"): st.session_state.mostrar_confirmacao_exclusao = None; st.rerun()
    
    if st.session_state.get('mostrar_form_colaborador'):
        st.markdown("---")
        st.subheader("📝 " + ("Editar Colaborador" if st.session_state.colaborador_editando else "Novo Colaborador"))
        colaborador = st.session_state.colaborador_editando or {}
        with st.form("form_colaborador"):
            col1, col2 = st.columns(2)
            with col1:
                colaborador_azure = st.text_input("Nome no Azure DevOps *", value=colaborador.get('colaborador_azure', ''))
                colaborador_ponto = st.text_input("Nome no Ponto *", value=colaborador.get('colaborador_ponto', ''))
                email = st.text_input("E-mail", value=colaborador.get('email', ''))
                cargo = st.text_input("Cargo *", value=colaborador.get('cargo', ''))
            with col2:
                nivel = st.selectbox("Nível *", options=["", "Senior", "Pleno", "Junior", "Portfólio"], index=["", "Senior", "Pleno", "Junior", "Portfólio"].index(colaborador.get('nivel', '')) if colaborador.get('nivel', '') in ["", "Senior", "Pleno", "Junior", "Portfólio"] else 0)
                status = st.selectbox("Status", options=["Ativo", "Inativo", "Bloqueado"], index=["Ativo", "Inativo", "Bloqueado"].index(colaborador.get('status', 'Ativo')) if colaborador.get('status', 'Ativo') in ["Ativo", "Inativo", "Bloqueado"] else 0)
                data_inicio = st.date_input("Data de Início *", value=colaborador.get('data_inicio') if colaborador.get('data_inicio') else None, format="DD/MM/YYYY")
                data_saida = st.date_input("Data de Saída", value=colaborador.get('data_saida') if colaborador.get('data_saida') else None, format="DD/MM/YYYY")
            col3, col4 = st.columns(2)
            with col3: promocao = st.date_input("Data da Promoção", value=colaborador.get('promocao') if colaborador.get('promocao') else None, format="DD/MM/YYYY")
            with col4: reposicao = st.text_input("Novo Colaborador (Reposição)", value=colaborador.get('reposicao', ''))
            
            st.markdown("---")
            col_submit, col_cancel = st.columns(2)
            with col_submit: submitted = st.form_submit_button("💾 Salvar", width="stretch", type="primary")
            with col_cancel: cancelled = st.form_submit_button("❌ Cancelar", width="stretch")
            if cancelled: st.session_state.colaborador_editando = None; st.session_state.mostrar_form_colaborador = False; st.rerun()
            if submitted:
                erros = []
                if not colaborador_azure.strip(): erros.append("Nome no Azure DevOps é obrigatório")
                if not colaborador_ponto.strip(): erros.append("Nome no Ponto é obrigatório")
                if not cargo.strip(): erros.append("Cargo é obrigatório")
                if not nivel: erros.append("Nível é obrigatório")
                if not data_inicio: erros.append("Data de Início é obrigatória")
                if erros:
                    for erro in erros: st.error(f"⚠️ {erro}")
                else:
                    dados = {'colaborador_azure': colaborador_azure, 'colaborador_ponto': colaborador_ponto, 'cargo': cargo, 'nivel': nivel, 'data_inicio': data_inicio if data_inicio else None, 'email': email, 'promocao': promocao if promocao else None, 'reposicao': reposicao, 'data_saida': data_saida if data_saida else None, 'status': status}
                    if st.session_state.colaborador_editando: sucesso, msg = atualizar_colaborador(st.session_state.colaborador_editando['id'], dados)
                    else: sucesso, msg = criar_colaborador(dados)
                    if sucesso: st.success(f"✅ {msg}"); st.session_state.colaborador_editando = None; st.session_state.mostrar_form_colaborador = False; st.rerun()
                    else: st.error(f"❌ {msg}")

# =============================================================================
# FATURAMENTO
# =============================================================================
@st.cache_data(ttl=120, show_spinner=False)
def verificar_faturamento_existente(ano: int, mes: int) -> Optional[Dict]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT id, ano, mes, qt_dias_uteis, total_faturado, criado_em FROM faturamento_mensal WHERE ano = %s AND mes = %s", (ano, mes))
            result = cur.fetchone()
        conn.close()
        return result
    except Exception as e:
        logger.error(f"Erro ao verificar faturamento existente: {e}")
        return None

@st.cache_data(ttl=120, show_spinner=False)
def buscar_detalhes_faturamento(id_faturamento: int) -> pd.DataFrame:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return pd.DataFrame()
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("SELECT profissional, cargo, nivel, data_inicio, promocao_cargo, novo_colaborador, data_saida, horas_trabalhadas, horas_afast_acordado, horas_afast_legal, horas_arredondadas, valor_profissional, valor_recalculado FROM faturamento_detalhe WHERE id_faturamento = %s ORDER BY profissional", (id_faturamento,))
            rows = cur.fetchall()
        conn.close()
        return pd.DataFrame(rows)
    except Exception as e:
        logger.error(f"Erro ao buscar detalhes do faturamento: {e}")
        return pd.DataFrame()

def salvar_faturamento_no_banco(ano: int, mes: int, dias_uteis: float, total_faturado: float, df_detalhes: pd.DataFrame) -> Tuple[bool, str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado."
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor(cursor_factory=RealDictCursor) as cur:
            cur.execute("INSERT INTO faturamento_mensal (ano, mes, qt_dias_uteis, total_faturado) VALUES (%s, %s, %s, %s) ON CONFLICT (ano, mes) DO UPDATE SET qt_dias_uteis = EXCLUDED.qt_dias_uteis, total_faturado = EXCLUDED.total_faturado, atualizado_em = CURRENT_TIMESTAMP RETURNING id", (int(ano), int(mes), float(dias_uteis), float(total_faturado)))
            id_faturamento = cur.fetchone()['id']
            cur.execute("DELETE FROM faturamento_detalhe WHERE id_faturamento = %s", (id_faturamento,))
            
            dados_detalhe = []
            for _, row in df_detalhes.iterrows():
                data_ini = datetime.strptime(str(row["Data de Início"]), "%d/%m/%Y").date() if row["Data de Início"] and row["Data de Início"] != "" else None
                data_saida = datetime.strptime(str(row["Data de Saída"]), "%d/%m/%Y").date() if row["Data de Saída"] and row["Data de Saída"] != "" else None
                dados_detalhe.append((int(id_faturamento), str(row["Profissional"]), str(row["Cargo"]), str(row["Nível"]), data_ini, str(row["Promoção de Cargo"]) if row["Promoção de Cargo"] else "", str(row["Novo Colaborador"]) if row["Novo Colaborador"] else "", data_saida, float(row["Horas Trabalhadas"]), float(row["Horas Afast Acordado"]), float(row["Horas Afast Legal"]), float(row["Horas Arredondadas"]), float(row["Valor Profissional"]), float(row["Valor Recalculado"])))
            
            query = """
                INSERT INTO faturamento_detalhe (id_faturamento, profissional, cargo, nivel, data_inicio, promocao_cargo, novo_colaborador, data_saida, horas_trabalhadas, horas_afast_acordado, horas_afast_legal, horas_arredondadas, valor_profissional, valor_recalculado)
                VALUES %s
            """
            execute_values(cur, query, dados_detalhe)
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return True, f"Faturamento salvo com sucesso! ID: {id_faturamento}"
    except Exception as e:
        if conn: conn.rollback(); conn.close()
        return False, f"Erro ao salvar: {str(e)}"

def deletar_faturamento_existente(ano: int, mes: int) -> Tuple[bool, str]:
    db_url = os.getenv("NEON_DATABASE_URL")
    if not db_url: return False, "Banco de dados não configurado."
    conn = None
    try:
        conn = psycopg2.connect(db_url, connect_timeout=10)
        with conn.cursor() as cur:
            cur.execute("DELETE FROM faturamento_mensal WHERE ano = %s AND mes = %s", (ano, mes))
            conn.commit()
        conn.close()
        st.cache_data.clear()
        return True, "Faturamento deletado com sucesso."
    except Exception as e:
        if conn: conn.rollback(); conn.close()
        return False, f"Erro ao deletar: {str(e)}"

def calcular_faturamento(df_colaboradores, df_horasponto, df_valores, dias_uteis):
    if 'status' in df_colaboradores.columns: df_colaboradores = df_colaboradores[df_colaboradores['status'] == 'Ativo'].copy()
    if df_colaboradores.empty: return pd.DataFrame()
    nivel_valor = {str(r["nivel"]).strip(): limpar_valor_monetario(r.get("valor")) for _, r in df_valores.iterrows() if str(r.get("nivel", "")).strip()}
    
    colaborador_horas = {}
    for _, row in df_horasponto.iterrows():
        ponto_name = str(row["colaborador_ponto"]).strip()
        horas_ponto = float(row["horas_ponto"]) if pd.notna(row.get("horas_ponto")) else 0.0
        afastamento = float(row["afastamento"]) if pd.notna(row.get("afastamento")) else 0.0
        if ponto_name: colaborador_horas[ponto_name] = {"horas_ponto": horas_ponto, "afastamento": afastamento}
    
    lista_faturamento = []
    for _, row in df_colaboradores.iterrows():
        colaborador_ponto = str(row["colaborador_ponto"]).strip()
        cargo, nivel = str(row.get("cargo", "")).strip(), str(row.get("nivel", "")).strip()
        data_inicio, data_saida = row.get("data_inicio"), row.get("data_saida")
        promocao = str(row.get("promocao", "")).strip() if pd.notna(row.get("promocao")) else ""
        reposicao = str(row.get("reposicao", "")).strip() if pd.notna(row.get("reposicao")) else ""
        horas_info = colaborador_horas.get(colaborador_ponto, {"horas_ponto": 0.0, "afastamento": 0.0})
        horas_trabalhadas, horas_afast_legal = horas_info["horas_ponto"], horas_info["afastamento"]
        horas_arredondadas = horas_trabalhadas
        valor_profissional = nivel_valor.get(nivel, 0.0)
        valor_diario = valor_profissional / dias_uteis if dias_uteis > 0 else 0.0
        valor_recalculado = round(horas_arredondadas * valor_diario, 2)
        lista_faturamento.append({"Profissional": colaborador_ponto, "Cargo": cargo, "Nível": nivel, "Data de Início": data_inicio.strftime("%d/%m/%Y") if pd.notna(data_inicio) else "", "Promoção de Cargo": promocao, "Novo Colaborador": reposicao, "Data de Saída": data_saida.strftime("%d/%m/%Y") if pd.notna(data_saida) else "", "Horas Trabalhadas": round(horas_trabalhadas, 2), "Horas Afast Acordado": 0.0, "Horas Afast Legal": round(horas_afast_legal, 2), "Horas Arredondadas": round(horas_arredondadas, 2), "Valor Profissional": round(valor_profissional, 2), "Valor Recalculado": valor_recalculado})
    return pd.DataFrame(lista_faturamento)

def render_faturamento_tab(df_colaboradores, df_horasponto, df_valores, dias_uteis, year, month_name, month_num):
    st.subheader("💰 Faturamento por Colaborador")
    st.info(f"Período: **{month_name}/{year}** | Dias úteis: **{dias_uteis/8:.0f}** ({dias_uteis:.0f}h)")
    faturamento_existente = verificar_faturamento_existente(year, month_num)
    if 'faturamento_calculado' not in st.session_state: st.session_state.faturamento_calculado = False
    if 'df_faturamento_exibido' not in st.session_state: st.session_state.df_faturamento_exibido = None
    if 'total_faturado_exibido' not in st.session_state: st.session_state.total_faturado_exibido = 0.0
    if 'mostrar_alerta_recuperar' not in st.session_state: st.session_state.mostrar_alerta_recuperar = False
    
    if faturamento_existente and not st.session_state.faturamento_calculado:
        st.info(f"✅ Faturamento existente encontrado para **{month_name}/{year}** (ID: {faturamento_existente['id']})")
        df_detalhes = buscar_detalhes_faturamento(faturamento_existente['id'])
        if not df_detalhes.empty:
            df_detalhes = df_detalhes.rename(columns={'profissional': 'Profissional', 'cargo': 'Cargo', 'nivel': 'Nível', 'data_inicio': 'Data de Início', 'promocao_cargo': 'Promoção de Cargo', 'novo_colaborador': 'Novo Colaborador', 'data_saida': 'Data de Saída', 'horas_trabalhadas': 'Horas Trabalhadas', 'horas_afast_acordado': 'Horas Afast Acordado', 'horas_afast_legal': 'Horas Afast Legal', 'horas_arredondadas': 'Horas Arredondadas', 'valor_profissional': 'Valor Profissional', 'valor_recalculado': 'Valor Recalculado'})
            for col in ['Data de Início', 'Data de Saída']:
                if col in df_detalhes.columns: df_detalhes[col] = df_detalhes[col].apply(lambda x: x.strftime("%d/%m/%Y") if pd.notna(x) else "")
            st.session_state.df_faturamento_exibido = df_detalhes
            st.session_state.total_faturado_exibido = faturamento_existente['total_faturado']
            st.session_state.faturamento_calculado = True
            st.rerun()
    
    st.markdown("---")
    col_btn1, col_btn2 = st.columns([3, 1])
    with col_btn1:
        if st.button("🧮 Calcular Faturamento", width="stretch", type="primary"):
            if faturamento_existente: st.session_state.mostrar_alerta_recuperar = True; st.rerun()
            else:
                with st.spinner("Calculando faturamento..."):
                    df_faturamento = calcular_faturamento(df_colaboradores, df_horasponto, df_valores, dias_uteis)
                    if df_faturamento.empty: st.warning("⚠️ Nenhum colaborador ativo para calcular faturamento.")
                    else:
                        total_faturado = df_faturamento["Valor Recalculado"].sum()
                        sucesso, msg = salvar_faturamento_no_banco(year, month_num, dias_uteis, total_faturado, df_faturamento)
                        if sucesso: st.session_state.df_faturamento_exibido = df_faturamento; st.session_state.total_faturado_exibido = total_faturado; st.session_state.faturamento_calculado = True; st.success(f"✅ {msg}"); st.rerun()
                        else: st.error(f"❌ {msg}")
    
    if st.session_state.get('mostrar_alerta_recuperar'):
        st.warning(f"⚠️ **Atenção:** Já existem dados de faturamento para **{month_name}/{year}** no banco de dados.")
        st.info(f"ID do faturamento existente: **{faturamento_existente['id']}** | Total: **{format_moeda_brl(faturamento_existente['total_faturado'])}**")
        col_sim, col_nao = st.columns(2)
        with col_sim:
            if st.button("✅ Sim, Recalcular", width="stretch", type="primary"):
                with st.spinner("Recalculando faturamento..."):
                    sucesso_del, msg_del = deletar_faturamento_existente(year, month_num)
                    if sucesso_del:
                        df_faturamento = calcular_faturamento(df_colaboradores, df_horasponto, df_valores, dias_uteis)
                        total_faturado = df_faturamento["Valor Recalculado"].sum()
                        sucesso_save, msg_save = salvar_faturamento_no_banco(year, month_num, dias_uteis, total_faturado, df_faturamento)
                        if sucesso_save:
                            st.session_state.df_faturamento_exibido = df_faturamento
                            st.session_state.total_faturado_exibido = total_faturado
                            st.session_state.faturamento_calculado = True
                            st.session_state.mostrar_alerta_recuperar = False
                            st.success(f"✅ Faturamento recalculado e salvo! {msg_save}")
                            st.rerun()
                        else: st.error(f"❌ {msg_save}")
                    else: st.error(f"❌ {msg_del}")
        with col_nao:
            if st.button("❌ Não, Cancelar", width="stretch"): st.session_state.mostrar_alerta_recuperar = False; st.rerun()
    
    if st.session_state.faturamento_calculado and st.session_state.df_faturamento_exibido is not None:
        df_faturamento, total_faturado = st.session_state.df_faturamento_exibido, st.session_state.total_faturado_exibido
        st.markdown("---")
        col_total1, col_total2 = st.columns([3, 1])
        with col_total1: st.markdown("### 💵 Total Faturado")
        with col_total2: st.metric(label="", value=format_moeda_brl(total_faturado))
        st.markdown("---")
        excel_data = generate_excel_faturamento(df_faturamento, month_name, year, dias_uteis, total_faturado)
        st.download_button("⬇️ Exportar Faturamento (Excel)", data=excel_data, file_name=f"faturamento_{year}_{month_name}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", width="stretch")
        st.markdown("---")
        st.dataframe(df_faturamento, width="stretch", hide_index=True)

# =============================================================================
# CRUD PERFIL
# =============================================================================
def render_crud_perfil():
    tipo_usuario, usuario_id = st.session_state.user_info.get('cargo', 'User'), st.session_state.user_info.get('id')
    st.markdown("---"); st.subheader("⚙️ Gerenciamento de Perfil")
    usuarios = listar_usuarios(tipo_usuario, usuario_id)
    if not usuarios: st.warning("Nenhum usuário encontrado."); return
    
    st.markdown("### Usuários Cadastrados")
    df_exibicao = pd.DataFrame(usuarios)
    df_exibicao['apikey'] = df_exibicao['apikey'].apply(lambda x: f"{x[:6]}...{x[-4:]}" if x and len(x) > 10 else x)
    df_exibicao['criado_em'] = pd.to_datetime(df_exibicao['criado_em']).dt.strftime('%d/%m/%Y %H:%M')
    st.dataframe(df_exibicao[['id', 'nome', 'email', 'tipo', 'status', 'apikey', 'criado_em']], width="stretch", hide_index=True)
    
    st.markdown("---")
    if tipo_usuario == 'Admin':
        usuario_selecionado = st.selectbox("Selecione o usuário para editar:", options=[f"{u['id']} - {u['nome']} ({u['email']})" for u in usuarios], index=0)
        usuario_id_selecionado = int(usuario_selecionado.split(" - ")[0])
        usuario_dados = next((u for u in usuarios if u['id'] == usuario_id_selecionado), None)
    else:
        usuario_dados, usuario_id_selecionado = usuarios[0], usuario_id
        st.info(f"Editando seu próprio perfil: **{usuario_dados['nome']}**")
    
    if not usuario_dados: return
    with st.form("form_editar_usuario"):
        nome = st.text_input("Nome Completo", value=usuario_dados.get('nome', ''))
        email = st.text_input("E-mail", value=usuario_dados.get('email', ''))
        apikey = st.text_input("API Key do Azure DevOps", value=usuario_dados.get('apikey', ''), type="password")
        senha_hash = st.text_input("Nova Senha (deixe em branco para manter)", type="password", placeholder="Mínimo 6 caracteres")
        if tipo_usuario == 'Admin':
            col1, col2 = st.columns(2)
            with col1: tipo = st.selectbox("Tipo de Usuário", options=['User', 'Admin', 'Gestor', 'Auditor'], index=['User', 'Admin', 'Gestor', 'Auditor'].index(usuario_dados.get('tipo', 'User')))
            with col2: status = st.selectbox("Status", options=['Ativo', 'Inativo', 'Bloqueado', 'Pendente'], index=['Ativo', 'Inativo', 'Bloqueado', 'Pendente'].index(usuario_dados.get('status', 'Ativo')))
        else: tipo, status = usuario_dados.get('tipo', 'User'), usuario_dados.get('status', 'Ativo')
        submitted = st.form_submit_button("💾 Salvar Alterações", width="stretch")
        if submitted:
            dados = {'nome': nome.strip(), 'email': email.lower().strip(), 'apikey': apikey.strip()}
            if senha_hash and len(senha_hash) >= 6: dados['senha_hash'] = senha_hash
            if tipo_usuario == 'Admin': dados['tipo'], dados['status'] = tipo, status
            sucesso, msg = atualizar_usuario(usuario_id_selecionado, dados, tipo_usuario, usuario_id)
            if sucesso:
                st.success(f"✅ {msg}")
                if usuario_id_selecionado == usuario_id:
                    st.session_state.user_info['nome_completo'] = nome.strip()
                    st.session_state.user_info['email'] = email.lower().strip()
                    if apikey.strip(): st.session_state.user_info['apikey'] = apikey.strip()
                st.rerun()
            else: st.error(f"❌ {msg}")

# =============================================================================
# LOGIN
# =============================================================================
def render_login_screen():
    inject_custom_css()
    st.markdown('<div class="dotted-rings"><div class="ring ring-1"></div><div class="ring ring-2"></div><div class="ring ring-3"></div><div class="ring ring-4"></div><div class="ring ring-5"></div></div>', unsafe_allow_html=True)
    col1, col2 = st.columns([1.1, 0.9], gap="large")
    with col1: st.markdown('<div style="margin-top: 40px;"><span class="badge-qween">Qween Platform</span><h1 class="hero-title">Relatório Mensal de acompanhamento da Execução</h1><p class="hero-subtitle">Secretaria Municipal de Educação - SP</p></div>', unsafe_allow_html=True)
    with col2:
        st.markdown('<div class="auth-card">', unsafe_allow_html=True)
        view = st.session_state.get("login_view", "login")
        if view == "login":
            st.markdown("<h2 style='color:white; font-weight:700;'>Entrar</h2>", unsafe_allow_html=True)
            st.markdown("<p style='color:#e6e6e6; margin-bottom:30px;'>Acesse com suas credenciais institucionais</p>", unsafe_allow_html=True)
            with st.form("login_form"):
                email = st.text_input("E-mail ou Usuário", placeholder="nome@spassu.com.br")
                password = st.text_input("Senha", type="password", placeholder="••••••••")
                submitted = st.form_submit_button("Acessar Painel", width="stretch")
                if submitted:
                    if email and password:
                        user = authenticate_user(email, password)
                        if user:
                            if user.get('status') != 'Ativo': st.error("⚠️ Usuário inativo. Contate o administrador.")
                            elif not user.get('apikey'): st.error("⚠️ Seu usuário não possui API Key do Azure DevOps cadastrada.")
                            else:
                                st.session_state.logged_in = True
                                st.session_state.user_info = {'id': user['id'], 'nome_completo': user['nome'], 'email': user['email'], 'cargo': user['tipo'], 'apikey': user['apikey']}
                                st.session_state.login_view = "login"
                                st.rerun()
                        else: st.error("E-mail ou senha inválidos.")
                    else: st.warning("Preencha todos os campos.")
            st.markdown("")
            col_btn1, col_btn2 = st.columns([3, 1])
            with col_btn2:
                if st.button("Esqueci minha senha", width="stretch"): st.session_state.login_view = "forgot"; st.rerun()
            st.markdown('<div style="display:flex; align-items:center; text-align:center; margin:28px 0; color:rgba(255,255,255,0.7); font-size:0.8rem; font-weight:600;"><div style="flex:1; border-bottom:1px solid rgba(255,255,255,0.25); margin-right:12px;"></div>OU<div style="flex:1; border-bottom:1px solid rgba(255,255,255,0.25); margin-left:12px;"></div></div>', unsafe_allow_html=True)
            if st.button("Primeiro Acesso? Ativar Conta", width="stretch"): st.session_state.login_view = "register"; st.rerun()
        elif view == "forgot":
            st.markdown("<h2 style='color:white; font-weight:700;'>Recuperação de Senha</h2>", unsafe_allow_html=True)
            st.markdown("<p style='color:#e6e6e6; margin-bottom:30px;'>Insira seu e-mail para receber o token de recuperação</p>", unsafe_allow_html=True)
            with st.form("forgot_form"):
                email = st.text_input("E-mail Cadastrado", placeholder="seu.email@spassu.com.br")
                submitted = st.form_submit_button("Gerar Token de Recuperação", width="stretch")
                if submitted:
                    if not email: st.error("⚠️ Informe o e-mail.")
                    else:
                        sucesso, msg, token = solicitar_reset_senha(email)
                        if sucesso:
                            st.success(f"✅ {msg}")
                            st.markdown(f'<div class="token-box">{token}</div>', unsafe_allow_html=True)
                            st.session_state.reset_token = token
                            st.session_state.reset_email = email
                            st.session_state.go_to_reset = True
                        else: st.error(f"❌ {msg}")
            if st.session_state.get('go_to_reset'):
                if st.button("Redefinir Senha Agora →", width="stretch"): st.session_state.login_view = "reset_password"; st.session_state.go_to_reset = False; st.rerun()
            if st.button("← Voltar ao Login", width="stretch"): st.session_state.login_view = "login"; st.session_state.go_to_reset = False; st.rerun()
        elif view == "reset_password":
            st.markdown("<h2 style='color:white; font-weight:700;'>Redefinir Senha</h2>", unsafe_allow_html=True)
            with st.form("reset_form"):
                token = st.text_input("Token de Recuperação", value=st.session_state.get('reset_token', ''))
                nova_senha = st.text_input("Nova Senha", type="password")
                confirmar_senha = st.text_input("Confirme a Nova Senha", type="password")
                submitted = st.form_submit_button("Redefinir Senha", width="stretch")
                if submitted:
                    if not token or not nova_senha or nova_senha != confirmar_senha: st.error("⚠️ Verifique os dados informados.")
                    else:
                        sucesso, msg = redefinir_senha(token, nova_senha)
                        if sucesso: st.success(f"✅ {msg}"); st.session_state.login_view = "login"; st.rerun()
                        else: st.error(f"❌ {msg}")
            if st.button("← Voltar ao Login", width="stretch"): st.session_state.login_view = "login"; st.rerun()
        elif view == "register":
            st.markdown("<h2 style='color:white; font-weight:700;'>Primeiro Acesso</h2>", unsafe_allow_html=True)
            with st.form("register_form", clear_on_submit=True):
                nome = st.text_input("Nome Completo")
                email = st.text_input("E-mail Corporativo")
                apikey = st.text_input("API Key do Azure DevOps")
                senha = st.text_input("Crie uma Senha", type="password")
                confirmar_senha = st.text_input("Confirme a Senha", type="password")
                submitted = st.form_submit_button("Validar e Criar Conta", width="stretch")
                if submitted:
                    if not nome or not email or not apikey or not senha or senha != confirmar_senha: st.error("⚠️ Preencha os campos corretamente.")
                    else:
                        sucesso, msg = criar_usuario_com_apikey(nome, email, senha, apikey.strip())
                        if sucesso: st.success(f"✅ {msg}"); st.session_state.login_view = "login"; st.rerun()
                        else: st.error(f"❌ {msg}")
            if st.button("Já possuo senha (Entrar)", width="stretch"): st.session_state.login_view = "login"; st.rerun()
        st.markdown('</div>', unsafe_allow_html=True)
    st.markdown(f'<div style="position:fixed;bottom:30px;right:40px;display:flex;align-items:center;gap:20px;z-index:100;"><img src="{SPASSU_LOGO_URL}" style="height:40px;"><div style="width:1px;height:40px;background:rgba(255,255,255,0.5);"></div><img src="{SME_LOGO_URL}" style="height:40px;"></div>', unsafe_allow_html=True)

# =============================================================================
# MAIN COM OTIMIZAÇÕES GERAIS DE FLUXO
# =============================================================================
def main():
    if 'logged_in' not in st.session_state: st.session_state.logged_in = False
    if 'login_view' not in st.session_state: st.session_state.login_view = "login"
    if 'show_perfil' not in st.session_state: st.session_state.show_perfil = False

    if not st.session_state.logged_in: render_login_screen(); return

    st.set_page_config(page_title="Dashboard de Entregas - SME Spassu", page_icon="📊", layout="wide", initial_sidebar_state="expanded")
    st.sidebar.markdown("---")
    st.sidebar.success(f"👤 {st.session_state.user_info['nome_completo']}")
    st.sidebar.caption(f"Tipo: {st.session_state.user_info.get('cargo', 'User')}")
    if st.sidebar.button("⚙️ Perfil", width="stretch"): st.session_state.show_perfil = True; st.rerun()
    if st.sidebar.button("🚪 Sair do Sistema", width="stretch"):
        st.session_state.logged_in = False
        st.session_state.user_info = None
        st.session_state.show_perfil = False
        st.rerun()

    if st.session_state.get('show_perfil'):
        st.title("⚙️ Gerenciamento de Perfil")
        render_crud_perfil()
        if st.button("← Voltar ao Dashboard"): st.session_state.show_perfil = False; st.rerun()
        return

    st.title("📊 Dashboard de Acompanhamento de Entregas")
    st.caption("Azure DevOps • SME-Spassu • Sprint Tracker")
    st.markdown("---")
    st.sidebar.header("⚙️ Configuração & Filtros")
    
    pat = st.session_state.user_info.get('apikey', '')
    if not pat: st.error("⚠️ API Key do Azure DevOps não cadastrada para este usuário."); st.stop()
    
    api_valida, msg_erro = validate_ado_api_key(pat)
    if not api_valida:
        st.error(f"🚫 {msg_erro}")
        st.stop()

    current_year = datetime.now().year
    year = st.sidebar.selectbox("📅 Ano", options=[current_year + 1, current_year, current_year - 1], index=1 if current_year == 2026 else 0)
    month_names = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]
    month_name = st.sidebar.selectbox("📅 Mês", options=month_names, index=datetime.now().month - 1)
    month_num = month_names.index(month_name) + 1

    all_projects = get_active_projects_from_db(pat)
    selected_projects = st.sidebar.multiselect("🏢 Projetos", options=all_projects, default=all_projects)
    if not selected_projects: st.info("Selecione ao menos um projeto."); st.stop()

    # Paralelização das chamadas de Sprints
    with st.spinner("🔄 Buscando sprints..."):
        with ThreadPoolExecutor(max_workers=5) as executor:
            futures = {executor.submit(fetch_all_sprints_for_project, pat, proj, year, month_num): proj for proj in selected_projects}
            sprints_by_project = {futures[f]: f.result() for f in as_completed(futures) if f.result()}
    if not sprints_by_project: st.warning(f"⚠️ Nenhuma sprint encontrada em {month_name}/{year}."); st.stop()

    sprint_options = [f"{s['path']}" for proj, sprints in sprints_by_project.items() for s in sprints]
    selected_sprints_display = st.sidebar.multiselect("🏃 Sprints", options=sprint_options, default=sprint_options)
    
    sprint_to_info = {}
    for proj, sprints in sprints_by_project.items():
        for s in sprints: sprint_to_info[f"{proj}|||{s['path']}"] = (proj, s["path"])
    selected_sprint_paths, paths_by_project = [], {}
    for display_name in selected_sprints_display:
        for key, (proj, path) in sprint_to_info.items():
            if path == display_name:
                selected_sprint_paths.append((proj, path))
                paths_by_project.setdefault(proj, []).append(path)
                break

    selected_wi_types = st.sidebar.multiselect("📋 Tipo de Work Item", options=ALL_WI_TYPES, default=DEFAULT_WI_TYPES)
    selected_states = st.sidebar.multiselect("🎯 Situação", options=ALL_STATES, default=ALL_STATES)

    # OTIMIZAÇÃO CRÍTICA: Busca paralela dos IDs de itens UMA ÚNICA VEZ
    with st.spinner("🔍 Buscando work items e hierarquias em paralelo..."):
        all_ids_by_project = {}
        children_map = {}
        
        with ThreadPoolExecutor(max_workers=6) as executor:
            fut_ids = {executor.submit(fetch_all_workitems_ids_by_project, pat, proj, tuple(paths)): proj for proj, paths in paths_by_project.items()}
            fut_links = {executor.submit(fetch_hierarchy_links_by_project, pat, proj, tuple(paths)): proj for proj, paths in paths_by_project.items()}
            
            for f in as_completed(fut_ids):
                proj = fut_ids[f]
                ids_res, _ = f.result()
                if ids_res: all_ids_by_project[proj] = ids_res
                
            for f in as_completed(fut_links):
                relations, _ = f.result()
                for rel in relations:
                    if rel and rel.get("source") and rel.get("target"):
                        src, tgt = rel["source"].get("id"), rel["target"].get("id")
                        if rel.get("rel") == "System.LinkTypes.Hierarchy-Forward": children_map.setdefault(src, []).append(tgt)
                        elif rel.get("rel") == "System.LinkTypes.Hierarchy-Reverse": children_map.setdefault(tgt, []).append(src)

        # Busca detalhes dos work items de forma unificada
        all_workitems, project_of_item = {}, {}
        with ThreadPoolExecutor(max_workers=5) as executor:
            fut_details = {executor.submit(fetch_workitems_details_batch, pat, proj, tuple(ids)): proj for proj, ids in all_ids_by_project.items() if ids}
            for f in as_completed(fut_details):
                proj = fut_details[f]
                items, _ = f.result()
                for item_id, item in items.items():
                    all_workitems[item_id] = item
                    project_of_item[item_id] = proj

    # Reutilização imediata em memória dos dados para Colaboradores (elimina 2ª requisição idêntica)
    all_workitems_colab = all_workitems
    project_of_item_colab = project_of_item

    # Filtragem em memória dos itens pais
    parents_data = []
    valid_paths = set(p for _, p in selected_sprint_paths)
    for item_id, item in all_workitems.items():
        f = item.get("fields", {})
        if f.get("System.WorkItemType") in selected_wi_types and f.get("System.State") in selected_states and f.get("System.IterationPath", "") in valid_paths:
            item["_project"] = project_of_item.get(item_id, "")
            parents_data.append(item)

    # Dados cadastrais rápidos via cache
    df_colaboradores = get_colaboradores_from_db()
    df_horasponto = get_horasponto_from_db(year, month_num)
    horas_uteis_mensais = get_dias_uteis_from_db(year, month_num)
    df_valores = get_valores_from_db()

    tipo_usuario = st.session_state.user_info.get('cargo', 'User')
    
    if tipo_usuario in ['Admin', 'Gestor']:
        tab_lista, tab_colaborador, tab_horas, tab_faturamento, tab_metricas, tab_ponto, tab_projetos, tab_crud_colab = st.tabs([
            "📋 Lista Geral", "👥 Tarefas por Colaborador", "🕐 Horas por Colaborador", "💰 Faturamento", "📊 Métricas", "🕐 Ponto", "🏢 Projetos", "⚙️ Colaboradores"
        ])
    elif tipo_usuario == 'Auditor':
        tab_lista, tab_colaborador, tab_horas, tab_faturamento, tab_metricas, tab_ponto = st.tabs([
            "📋 Lista Geral", "👥 Tarefas por Colaborador", "🕐 Horas por Colaborador", "💰 Faturamento", "📊 Métricas", "🕐 Ponto"
        ])
    else:
        tab_lista, tab_colaborador, tab_horas, tab_metricas, tab_ponto = st.tabs([
            "📋 Lista Geral", "👥 Tarefas por Colaborador", "🕐 Horas por Colaborador", "📊 Métricas", "🕐 Ponto"
        ])
    
    with tab_lista:
        total_items = len(parents_data)
        total_hours = sum(float(p.get("fields", {}).get("Microsoft.VSTS.Scheduling.OriginalEstimate") or 0) for p in parents_data)
        done_resolved = sum(1 for p in parents_data if p.get("fields", {}).get("System.State") in ["Done", "Resolved"])
        in_progress_todo = sum(1 for p in parents_data if p.get("fields", {}).get("System.State") in ["In Progress", "To Do"])
        historias_nao_done = sum(1 for p in parents_data if p.get("fields", {}).get("System.WorkItemType") == "Product Backlog Item" and p.get("fields", {}).get("System.State") != "Done")
        
        col1, col2, col3, col4 = st.columns(4)
        col1.metric("📋 Total de Itens", total_items)
        col2.metric("⏱️ Total de Horas", f"{total_hours:.2f} h")
        col3.metric("📊 Situação", f"✅ {done_resolved} |  {in_progress_todo}")
        col4.metric("📖 Histórias em Aberto", historias_nao_done, help="Product Backlog Items com status diferente de Done")
        st.markdown("---")
        
        if parents_data:
            type_counts = {}
            for p in parents_data:
                t = p.get("fields", {}).get("System.WorkItemType", "Desconhecido")
                type_counts[t] = type_counts.get(t, 0) + 1
            df_types = pd.DataFrame([{"Tipo": k, "Quantidade": v} for k, v in sorted(type_counts.items(), key=lambda x: x[1], reverse=True)])
            
            state_counts = {}
            for p in parents_data:
                s = p.get("fields", {}).get("System.State", "Desconhecido")
                state_counts[s] = state_counts.get(s, 0) + 1
            df_states = pd.DataFrame([{"Situação": k, "Quantidade": v} for k, v in sorted(state_counts.items(), key=lambda x: x[1], reverse=True)])
            
            state_colors = {"Done": "#28a745", "Resolved": "#17a2b8", "In Progress": "#ffc107", "To Do": "#6c757d"}
            col_chart1, col_chart2 = st.columns(2)
            with col_chart1:
                st.subheader("📊 Total por Tipo")
                fig_type = px.pie(df_types, values="Quantidade", names="Tipo", hole=0.5, color_discrete_sequence=px.colors.qualitative.Set2)
                fig_type.update_layout(margin=dict(t=20, b=20, l=20, r=20), height=350, showlegend=True)
                st.plotly_chart(fig_type, width="stretch")
            with col_chart2:
                st.subheader("📊 Total por Situação")
                fig_state = px.pie(df_states, values="Quantidade", names="Situação", hole=0.5, color="Situação", color_discrete_map=state_colors)
                fig_state.update_layout(margin=dict(t=20, b=20, l=20, r=20), height=350, showlegend=True)
                st.plotly_chart(fig_state, width="stretch")
            
            st.markdown("---")
            st.subheader("📥 Exportar Dados")
            col_csv1, col_csv2 = st.columns(2)
            with col_csv1:
                csv_parents = generate_csv_parents_only(parents_data, children_map, all_workitems)
                st.download_button("⬇️ Baixar CSV (Apenas Pais)", data=csv_parents, file_name=f"sme_sprints_{year}_{month_name}_pais.csv", mime="text/csv", width="stretch")
            with col_csv2:
                csv_complete = generate_csv_complete(parents_data, children_map, all_workitems)
                st.download_button("⬇️ Baixar CSV (Completo)", data=csv_complete, file_name=f"sme_sprints_{year}_{month_name}_completo.csv", mime="text/csv", width="stretch")
            
            render_table_html(parents_data, children_map, all_workitems, total_hours)
        else:
            st.warning("⚠️ Nenhum item após filtros.")
    
    with tab_colaborador:
        st.subheader("👥 Tarefas por Colaborador")
        st.info(f"Período: **{month_name}/{year}** | Apenas **State = Done**")
        if not df_colaboradores.empty: render_colaborador_tab(all_workitems_colab, project_of_item_colab, df_colaboradores, year, month_name)
        else: st.warning("⚠️ Tabela de colaboradores vazia.")
    
    with tab_horas:
        st.subheader("🕐 Horas por Colaborador")
        if df_colaboradores.empty: st.warning("⚠️ Tabela de colaboradores vazia.")
        elif df_horasponto.empty: st.warning(f"⚠️ Sem dados em `horasponto` para {month_name}/{year}.")
        elif horas_uteis_mensais == 0: st.warning(f"⚠️ Sem dados em `dias_uteis_mensais` para {month_name}/{year}.")
        else: render_horas_colaborador_tab(df_colaboradores, df_horasponto, horas_uteis_mensais, all_workitems_colab, project_of_item_colab, year, month_name)
    
    if tipo_usuario in ['Admin', 'Gestor', 'Auditor']:
        with tab_faturamento:
            if df_colaboradores.empty: st.warning("⚠️ Tabela de colaboradores vazia.")
            elif df_horasponto.empty: st.warning(f"⚠️ Sem dados em `horasponto` para {month_name}/{year}.")
            elif df_valores.empty: st.warning("⚠️ Tabela de valores vazia.")
            else: render_faturamento_tab(df_colaboradores, df_horasponto, df_valores, horas_uteis_mensais, year, month_name, month_num)
    
    with tab_metricas:
        render_metricas_tab(all_workitems, project_of_item, df_colaboradores, df_horasponto, df_valores, horas_uteis_mensais, year, month_name, month_num, selected_wi_types, selected_states, parents_data)
    
    with tab_ponto:
        render_ponto_tab(year, month_name, month_num, tipo_usuario)
    
    if tipo_usuario in ['Admin', 'Gestor']:
        with tab_projetos:
            render_projetos_tab()
    
    if tipo_usuario in ['Admin', 'Gestor']:
        with tab_crud_colab:
            render_colaboradores_tab()

if __name__ == "__main__":
    main()
